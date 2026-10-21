"""
Pruebas de la API y de los reportes. Corren sin las políticas de RLS (las aplica
el entrypoint con psql, no las migraciones), así que cubren la lógica de la
aplicación y los permisos de Django, no el aislamiento por parroquia.

    python manage.py test
"""

from datetime import date, timedelta
from decimal import Decimal
from unittest import mock

import tablib
from django.contrib.auth.models import Permission, User
from django.core.cache import cache
from django.core.management import call_command
from django.test import TestCase
from django.utils import timezone

from . import tasas
from .models import (
    AreaServicio, AsignacionRol, AsignacionServicio, Charla, Ciudad, Diocesis, Estado,
    EstadoPago, Gasto, Inscripcion, Pago, Pais, Parroquia, Participacion, Persona, Retiro, Rol,
    RolSistema, Tarea, TipoRetiro, Usuario,
)
from .reportes import PersonaResource


def rol_de(codigo):
    """El Rol del sistema con ese código (los carga la migración 0014)."""
    return Rol.objects.get(codigo=codigo)


class Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser("admin", "a@example.com", "x")
        cls.parroquia = Parroquia.objects.create(
            nombre="San José", diocesis=Diocesis.objects.create(nombre="Diócesis")
        )
        cls.retiro = Retiro.objects.create(
            parroquia=cls.parroquia, tipo=TipoRetiro.objects.create(nombre="Emaús hombres"),
            edicion=12, fecha_inicio=date(2027, 3, 12), fecha_fin=date(2027, 3, 14),
            cupo=3, cuota=Decimal("100"),
        )

    def setUp(self):
        self.client.force_login(self.admin)
        # Las pruebas no salen a internet: el BCV "publica" 180 Bs por dólar y
        # 200 por euro, salvo que una prueba cambie self.bcv.side_effect.
        cache.clear()
        self.bcv = mock.patch.object(tasas, "_consultar", return_value={
            "bolivares": {"VES": Decimal(1), "USD": Decimal(180), "EUR": Decimal(200)},
            "fecha": "2027-03-01",
        }).start()
        self.addCleanup(mock.patch.stopall)

    def persona(self, nombre, **datos):
        return Persona.objects.create(parroquia=self.parroquia, nombre=nombre, **datos)

    def post(self, ruta, datos):
        return self.client.post(ruta, datos, content_type="application/json")


class EstadoPagoTests(TestCase):
    def test_segun_lo_pagado(self):
        self.assertEqual(EstadoPago.de(None, 100), EstadoPago.PENDIENTE)
        self.assertEqual(EstadoPago.de(Decimal("40"), 100), EstadoPago.PARCIAL)
        self.assertEqual(EstadoPago.de(Decimal("100"), 100), EstadoPago.PAGADO)
        self.assertEqual(EstadoPago.de(None, 0), EstadoPago.PAGADO)  # sin cuota no debe


class ModulosTests(Base):
    def test_exige_sesion(self):
        self.client.logout()
        for ruta in ("/api/retiros/", "/api/esquema/", "/api/resumen/",
                     f"/api/retiros/{self.retiro.pk}/panel/"):
            # 401 (no 403): el frontend sabe que es la sesión, no el rol.
            self.assertEqual(self.client.get(ruta).status_code, 401, ruta)

    def test_esquema_describe_los_modulos(self):
        modulos = {m["slug"]: m for m in self.client.get("/api/esquema/").json()}
        self.assertIn("pagos", modulos)
        inscripciones = modulos["inscripciones"]
        pago = next(c for c in inscripciones["lista"] if c["nombre"] == "estado_pago")
        self.assertEqual(pago["tipo"], "opcion")
        self.assertEqual(len(pago["opciones"]), 3)
        self.assertTrue(inscripciones["puede_agregar"])

    def test_sin_permiso_de_pagos_no_ve_columnas_de_pago(self):
        user = User.objects.create_user("secretaria", password="x")
        user.user_permissions.add(Permission.objects.get(codename="view_inscripcion"))
        self.client.force_login(user)
        Inscripcion.objects.create(retiro=self.retiro, persona=self.persona("Ana"))

        (modulo,) = self.client.get("/api/esquema/").json()
        self.assertNotIn("estado_pago", [c["nombre"] for c in modulo["lista"]])
        self.assertFalse(modulo["puede_agregar"])
        (fila,) = self.client.get("/api/inscripciones/").json()["results"]
        self.assertNotIn("total_pagado", fila)
        self.assertNotIn("salud_alimentacion", fila)  # sin rol de RLS: nada confidencial
        self.assertEqual(self.post("/api/inscripciones/", {}).status_code, 403)

    def test_lista_de_inscripciones_trae_el_estado_de_pago(self):
        inscripcion = Inscripcion.objects.create(retiro=self.retiro, persona=self.persona("Ana"))
        Pago.objects.create(retiro=self.retiro, inscripcion=inscripcion, monto=40,
                            concepto=Pago.Concepto.CUOTA, fecha=date(2027, 1, 5))
        (fila,) = self.client.get("/api/inscripciones/").json()["results"]
        self.assertEqual(fila["estado_pago"], "parcial")
        self.assertEqual(fila["total_pagado"], 40)
        self.assertEqual(fila["persona_etiqueta"], "Ana")

    def test_validaciones_del_modelo_salen_junto_al_campo(self):
        r = self.client.patch(
            f"/api/retiros/{self.retiro.pk}/", {"fecha_fin": "2027-03-01"},
            content_type="application/json",
        )
        self.assertEqual(r.status_code, 400)
        self.assertIn("fecha_fin", r.json())

        ana = self.persona("Ana")
        r = self.post("/api/inscripciones/", {
            "retiro": str(self.retiro.pk), "persona": str(ana.pk), "padrino": str(ana.pk),
        })
        self.assertIn("padrino", r.json())

        r = self.post("/api/gastos/", {
            "retiro": str(self.retiro.pk), "categoria": "comida", "descripcion": "Pan",
            "monto": 0, "fecha": "2027-03-01",
        })
        self.assertIn("monto", r.json())

    def pago(self, **datos):
        return self.post("/api/pagos/", {
            "retiro": str(self.retiro.pk), "concepto": "donativo", "fecha": "2027-03-01",
            **datos,
        })

    def test_la_cuota_se_liga_a_la_inscripcion_con_solo_elegir_la_persona(self):
        ana = self.persona("Ana")
        inscripcion = Inscripcion.objects.create(retiro=self.retiro, persona=ana)
        r = self.pago(monto=40, concepto="cuota", persona=str(ana.pk))
        self.assertEqual(r.json()["inscripcion"], str(inscripcion.pk))

    def test_pago_en_la_moneda_base(self):
        # Sin moneda toma la base, y la tasa sobra aunque la manden.
        r = self.pago(monto=50, tasa=180)
        self.assertEqual(r.status_code, 201, r.content)
        pago = r.json()
        self.assertEqual((pago["moneda"], pago["tasa"], pago["monto_base"]), ("USD", None, 50))

    def test_pago_en_otra_moneda_se_convierte_con_su_tasa(self):
        self.assertIn("moneda", self.pago(monto=10, moneda="MXN").json())

        r = self.pago(monto=9000, moneda="VES", tasa="180.5")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.json()["monto_base"], 49.86)
        self.bcv.assert_not_called()  # con la tasa escrita no se consulta el BCV

        # Al corregir la tasa se recalcula el equivalente.
        r = self.client.patch(f"/api/pagos/{r.json()['id']}/", {"tasa": 180},
                              content_type="application/json")
        self.assertEqual(r.json()["monto_base"], 50)

    def test_sin_tasa_se_usa_la_del_bcv(self):
        pago = self.pago(monto=9000, moneda="VES").json()
        self.assertEqual((pago["tasa"], pago["monto_base"]), (180, 50))
        # Euros a dólares pasando por bolívares: 180 / 200 = 0,9 EUR por dólar.
        pago = self.pago(monto=45, moneda="EUR").json()
        self.assertEqual((pago["tasa"], pago["monto_base"]), (0.9, 50))
        self.assertEqual(self.bcv.call_count, 1)  # la segunda sale de la caché

        # El pago conserva su tasa aunque el BCV publique otra después.
        cache.clear()
        self.bcv.return_value["bolivares"]["USD"] = Decimal(190)
        r = self.client.patch(f"/api/pagos/{pago['id']}/", {"referencia": "A-1"},
                              content_type="application/json")
        self.assertEqual(r.json()["tasa"], 0.9)

        tasas_bcv = self.client.get("/api/resumen/").json()["tasas_bcv"]
        self.assertEqual(tasas_bcv, {"fecha": "2027-03-01", "USD": 190, "EUR": 200})

    def test_si_el_bcv_no_responde_se_pide_la_tasa(self):
        self.bcv.side_effect = OSError("sin red")
        with self.assertLogs("retiros.tasas", "ERROR"):
            r = self.pago(monto=9000, moneda="VES")
        self.assertEqual(r.status_code, 400)
        self.assertIn("tasa", r.json())
        self.assertEqual(self.pago(monto=9000, moneda="VES", tasa=180).status_code, 201)
        self.assertEqual(self.pago(monto=20).status_code, 201)  # en dólares no hace falta
        self.assertIsNone(self.client.get("/api/resumen/").json()["tasas_bcv"])
        self.assertEqual(self.bcv.call_count, 1)  # el fallo no se reintenta en cada petición

    def test_la_cuota_se_cubre_con_pagos_en_varias_monedas(self):
        inscripcion = Inscripcion.objects.create(retiro=self.retiro, persona=self.persona("Ana"))
        comunes = {"inscripcion": str(inscripcion.pk), "concepto": "cuota"}
        self.pago(monto=9000, moneda="VES", tasa=180, **comunes)   # 50 USD
        self.pago(monto=46, moneda="EUR", tasa="0.92", **comunes)  # 50 USD
        (fila,) = self.client.get("/api/inscripciones/").json()["results"]
        self.assertEqual((fila["total_pagado"], fila["estado_pago"]), (100, "pagado"))
        self.assertEqual(self.client.get("/api/resumen/").json()["ingresos"], 100)

    def test_inscripcion_duplicada_y_borrado_protegido(self):
        ana = self.persona("Ana")
        datos = {"retiro": str(self.retiro.pk), "persona": str(ana.pk)}
        self.assertEqual(self.post("/api/inscripciones/", datos).status_code, 201)
        r = self.post("/api/inscripciones/", datos)
        self.assertEqual(r.json()["persona"], ["Ya está inscrita en este retiro."])
        self.assertEqual(self.client.delete(f"/api/personas/{ana.pk}/").status_code, 409)

    def test_sesion_informa_moneda_y_region(self):
        sesion = self.client.get("/api/auth/sesion/").json()
        self.assertEqual((sesion["moneda"], sesion["region"]), ("USD", "es-VE"))


class PanelRetiroTests(Base):
    def setUp(self):
        super().setUp()
        ana, luis, marta, pedro = (self.persona(n) for n in ("Ana", "Luis", "Marta", "Pedro"))
        pagada = Inscripcion.objects.create(
            retiro=self.retiro, persona=ana, padrino=luis,
            estado=Inscripcion.Estado.CONFIRMADA,
        )
        Inscripcion.objects.create(retiro=self.retiro, persona=marta)
        Inscripcion.objects.create(
            retiro=self.retiro, persona=pedro, estado=Inscripcion.Estado.CANCELADA
        )
        Pago.objects.create(retiro=self.retiro, inscripcion=pagada, persona=ana, monto=100,
                            concepto=Pago.Concepto.CUOTA, fecha=date(2027, 1, 5))
        Pago.objects.create(retiro=self.retiro, monto=30, concepto=Pago.Concepto.DONATIVO,
                            fecha=date(2027, 1, 6))
        Gasto.objects.create(retiro=self.retiro, categoria=Gasto.Categoria.COMIDA,
                             descripcion="Mercado", monto=50, fecha=date(2027, 3, 1))
        cocina = AreaServicio.objects.create(nombre="Cocina")
        servidor = Participacion.objects.create(
            retiro=self.retiro, persona=luis, tipo=Participacion.Tipo.SERVIDOR
        )
        AsignacionServicio.objects.create(participacion=servidor, area=cocina,
                                          es_coordinador=True)
        Participacion.objects.create(
            retiro=self.retiro, persona=pedro, tipo=Participacion.Tipo.SERVIDOR
        )
        Charla.objects.create(retiro=self.retiro, tema="El hijo pródigo", orden=1,
                              expositor=luis)
        self.ruta = f"/api/retiros/{self.retiro.pk}/panel/"

    def test_junta_caminantes_equipo_charlas_y_finanzas(self):
        panel = self.client.get(self.ruta).json()

        caminantes = panel["caminantes"]
        self.assertEqual((caminantes["inscritos"], caminantes["disponibles"]), (2, 1))
        ana = caminantes["filas"][0]
        self.assertEqual((ana["persona"], ana["padrino"], ana["estado_pago"], ana["saldo"]),
                         ("Ana", "Luis", "pagado", 0))
        self.assertEqual(caminantes["filas"][1]["estado_pago"], "pendiente")

        equipo = panel["equipo"]
        self.assertEqual(equipo["total"], 2)
        self.assertEqual(equipo["areas"], [
            {"area": "Cocina", "personas": [{"nombre": "Luis", "es_coordinador": True}]}
        ])
        self.assertEqual(equipo["sin_area"], ["Pedro"])

        self.assertEqual(panel["charlas"][0]["expositor"], "Luis")

        finanzas = panel["finanzas"]
        self.assertEqual((finanzas["ingresos"], finanzas["egresos"], finanzas["balance"]),
                         (130, 50, 80))
        # Falta la cuota de Marta; la inscripción cancelada de Pedro no cuenta.
        self.assertEqual(finanzas["por_cobrar"], 100)
        self.assertEqual(
            {c["valor"]: c["total"] for c in finanzas["por_concepto"] if c["total"]},
            {"cuota": 100, "donativo": 30},
        )

    def test_ficha_de_cada_caminante(self):
        ana = Inscripcion.objects.get(persona__nombre="Ana")
        aragua = Estado.objects.create(
            pais=Pais.objects.create(nombre="Venezuela", codigo="VE"), nombre="Aragua")
        ana.persona.fecha_nacimiento = date(1990, 3, 13)
        ana.persona.estado = aragua
        ana.persona.ciudad = Ciudad.objects.create(estado=aragua, nombre="Maracay")
        ana.persona.documento = "12345678"
        ana.persona.datos_sensibles = {"telefono_casa": "0243-5551234"}
        ana.persona.save()
        ana.salud_alimentacion = {"tratamiento": "Insulina", "regimen_alimenticio": "Sin sal",
                                  "requerimientos": "Silla de ruedas"}
        ana.contacto_emergencia = {"nombre": "Rosa", "relacion": "Madre", "telefono": "0412",
                                   "celular": "0424", "nombre2": "Juan", "telefono2": "0416",
                                   "nombre5": "Luz", "relacion5": "Tía", "telefono5": "0414"}
        ana.save()

        ficha = self.client.get(self.ruta).json()["caminantes"]["filas"][0]["ficha"]
        self.assertEqual(ficha["edad"], 36)  # cumple 37 el día después de empezar
        self.assertEqual(
            (ficha["procedencia"], ficha["documento"], ficha["telefono_casa"],
             ficha["tratamiento"], ficha["alimentacion"], ficha["requerimientos"]),
            ("Maracay, Aragua", "V-12345678", "0243-5551234", "Insulina", "Sin sal",
             "Silla de ruedas"),
        )
        self.assertEqual(ficha["contactos"], [
            {"nombre": "Rosa", "relacion": "Madre", "telefono": "0412 / 0424"},
            {"nombre": "Juan", "relacion": "", "telefono": "0416"},
            {"nombre": "Luz", "relacion": "Tía", "telefono": "0414"},
        ])

    def test_lo_que_debe_cada_servidor(self):
        self.retiro.cuota_servidor = 50
        self.retiro.save()
        luis = Persona.objects.get(nombre="Luis")
        Pago.objects.create(retiro=self.retiro, persona=luis, monto=20,
                            concepto=Pago.Concepto.APORTACION, fecha=date(2027, 1, 7))

        servidores = self.client.get(self.ruta).json()["servidores"]
        self.assertEqual(
            [(f["persona"], f["areas"], f["pagado"], f["saldo"], f["estado_pago"])
             for f in servidores["filas"]],
            [("Luis", "Cocina", 20, 30, "parcial"), ("Pedro", "", 0, 50, "pendiente")],
        )
        self.assertEqual(servidores["por_cobrar"], 80)

    def test_cada_seccion_exige_su_permiso(self):
        user = User.objects.create_user("servidor", password="x")
        self.client.force_login(user)
        self.assertEqual(self.client.get(self.ruta).status_code, 403)

        user.user_permissions.add(
            *Permission.objects.filter(codename__in=["view_retiro", "view_inscripcion"])
        )
        panel = self.client.get(self.ruta).json()
        self.assertIsNone(panel["finanzas"])
        self.assertIsNone(panel["equipo"])
        self.assertIsNone(panel["servidores"])
        tablero = self.client.get("/api/resumen/").json()["tablero"]
        self.assertIsNone(tablero["finanzas"])
        self.assertIsNone(tablero["tareas"])
        self.assertNotIn("pagos", tablero["caminantes"])
        self.assertNotIn("pagado", panel["caminantes"]["filas"][0])
        self.assertNotIn("ficha", panel["caminantes"]["filas"][0])  # sin rol confidencial

    def test_tablero_del_inicio(self):
        hoy = timezone.localdate()
        Persona.objects.filter(nombre="Ana").update(
            fecha_nacimiento=(hoy + timedelta(days=5)).replace(year=1990))
        Persona.objects.filter(nombre="Luis").update(fecha_nacimiento=date(1980, 3, 13))
        Tarea.objects.create(retiro=self.retiro, titulo="Comprar velas",
                             fecha_limite=hoy - timedelta(days=1))
        Tarea.objects.create(retiro=self.retiro, titulo="Reservar bus")
        Tarea.objects.create(retiro=self.retiro, titulo="Hecha", hecha=True)

        tablero = self.client.get("/api/resumen/").json()["tablero"]
        self.assertEqual(tablero["retiro"]["id"], str(self.retiro.pk))
        self.assertEqual(tablero["dias_para_inicio"], (self.retiro.fecha_inicio - hoy).days)

        caminantes = tablero["caminantes"]
        self.assertEqual(
            (caminantes["inscritos"], caminantes["confirmados"], caminantes["por_confirmar"],
             caminantes["cancelados"], caminantes["por_cobrar"]),
            (2, 1, 1, 1, 100),
        )
        self.assertEqual({p["valor"]: p["total"] for p in caminantes["pagos"]},
                         {"pagado": 1, "parcial": 0, "pendiente": 1})
        # Ana (confirmada) pagó su cuota completa.
        self.assertEqual(caminantes["confirmados_pagados"], 1)
        self.assertEqual(caminantes["atencion"],
                         {"activos": 2, "salud": 0, "comida": 0, "requerimientos": 0})
        self.assertEqual(tablero["servidores"]["por_area"], [{"area": "Cocina", "total": 1}])

        tareas = tablero["tareas"]
        self.assertEqual((tareas["pendientes"], tareas["vencidas"], tareas["hechas"]), (2, 1, 1))
        self.assertEqual([t["titulo"] for t in tareas["proximas"]],
                         ["Comprar velas", "Reservar bus"])

        self.assertEqual(
            [(c["nombre"], c["durante_retiro"]) for c in tablero["cumpleanos"]],
            [("Ana", False), ("Luis", True)],
        )
        alertas = {a["clave"]: a["total"] for a in tablero["alertas"]}
        self.assertEqual(alertas["por_confirmar"], 1)
        self.assertEqual(alertas["caminantes_deben"], 1)
        self.assertEqual(alertas["sin_area"], 1)
        self.assertEqual(
            [p["total"] for p in tablero["finanzas"]["recaudacion"]], [100, 130])

    def test_tablero_de_otro_retiro(self):
        otro = Retiro.objects.create(
            parroquia=self.parroquia, tipo=self.retiro.tipo, edicion=13,
            fecha_inicio=date(2027, 9, 1), fecha_fin=date(2027, 9, 3),
        )
        resumen = self.client.get(f"/api/resumen/?retiro={otro.pk}").json()
        self.assertEqual(resumen["tablero"]["retiro"]["id"], str(otro.pk))
        self.assertEqual(len(resumen["retiros_tablero"]), 2)

    def test_retiro_inexistente(self):
        ruta = "/api/retiros/00000000-0000-0000-0000-000000000000/panel/"
        self.assertEqual(self.client.get(ruta).status_code, 404)

    def test_resumen(self):
        resumen = self.client.get("/api/resumen/").json()
        self.assertEqual((resumen["inscripciones"], resumen["balance"]), (3, 80))
        self.assertEqual(resumen["proximos_retiros"][0]["inscritos"], 2)


class AccesoPorRolTests(Base):
    """Coordinación y logística entran con su cuenta y obtienen permisos de su rol."""

    def entrar_como(self, rol):
        cuenta = User.objects.create_user(rol, password="x")
        usuario = Usuario.objects.create(persona=self.persona(rol), cuenta=cuenta)
        AsignacionRol.objects.create(usuario=usuario, rol=rol_de(rol), retiro=self.retiro)
        self.client.force_login(cuenta)
        return usuario

    def setUp(self):
        super().setUp()
        self.ana = Inscripcion.objects.create(retiro=self.retiro, persona=self.persona("Ana"))
        self.ruta = f"/api/retiros/{self.retiro.pk}/panel/"

    def test_sin_rol_no_ve_nada(self):
        self.client.force_login(User.objects.create_user("nadie", password="x"))
        self.assertEqual(self.client.get(self.ruta).status_code, 403)
        self.assertEqual(self.client.get("/api/esquema/").json(), [])

    def test_logistica_edita_caminantes_y_consulta_pagos(self):
        self.entrar_como(RolSistema.LOGISTICA)
        panel = self.client.get(self.ruta).json()
        self.assertIn("ficha", panel["caminantes"]["filas"][0])
        self.assertIsNotNone(panel["finanzas"])

        r = self.client.patch(f"/api/inscripciones/{self.ana.pk}/",
                              {"salud_alimentacion": {"tratamiento": "Insulina"}},
                              content_type="application/json")
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(self.client.get("/api/pagos/").status_code, 200)
        pago = {"retiro": str(self.retiro.pk), "concepto": "cuota", "monto": 10,
                "fecha": "2027-03-01", "inscripcion": str(self.ana.pk)}
        self.assertEqual(self.post("/api/pagos/", pago).status_code, 403)

    def test_coordinacion_registra_pagos(self):
        self.entrar_como(RolSistema.COORD_RETIRO)
        pago = {"retiro": str(self.retiro.pk), "concepto": "cuota", "monto": 100,
                "fecha": "2027-03-01", "inscripcion": str(self.ana.pk)}
        self.assertEqual(self.post("/api/pagos/", pago).status_code, 201)
        (fila,) = self.client.get(self.ruta).json()["caminantes"]["filas"]
        self.assertEqual((fila["estado_pago"], fila["saldo"]), ("pagado", 0))

    def test_el_rol_fija_parroquia_y_tipo_del_listado_de_retiros(self):
        samuel = TipoRetiro.objects.create(nombre="Samuel")
        otro = Retiro.objects.create(
            parroquia=self.parroquia, tipo=samuel, edicion=1,
            fecha_inicio=date(2027, 5, 1), fecha_fin=date(2027, 5, 3),
        )
        self.entrar_como(RolSistema.COORD_RETIRO)

        retiros = self.client.get("/api/retiros/").json()["results"]
        self.assertEqual([r["id"] for r in retiros], [str(self.retiro.pk)])
        # Pedir el otro tipo por la URL no lo muestra.
        r = self.client.get(f"/api/retiros/?tipo={samuel.pk}").json()
        self.assertEqual([x["id"] for x in r["results"]], [str(self.retiro.pk)])
        self.assertNotEqual(otro.pk, self.retiro.pk)

        (modulo,) = [m for m in self.client.get("/api/esquema/").json()
                     if m["slug"] == "retiros"]
        fijos = {f["nombre"]: f.get("fijo") for f in modulo["filtros"]}
        self.assertEqual(fijos["parroquia"],
                         {"valor": str(self.parroquia.pk), "etiqueta": "San José"})
        self.assertEqual(fijos["tipo"]["etiqueta"], "Emaús hombres")
        self.assertIsNone(fijos["estado"])

        # El inicio y el panel tampoco muestran el retiro del otro tipo.
        resumen = self.client.get("/api/resumen/").json()
        self.assertEqual([r["id"] for r in resumen["retiros_tablero"]], [str(self.retiro.pk)])
        self.assertEqual(self.client.get(f"/api/retiros/{otro.pk}/panel/").status_code, 404)

    def test_superusuario_no_tiene_filtros_fijos(self):
        (modulo,) = [m for m in self.client.get("/api/esquema/").json()
                     if m["slug"] == "retiros"]
        self.assertFalse(any("fijo" in f for f in modulo["filtros"]))

    def test_tamano_de_pagina(self):
        for n in range(3):
            self.persona(f"Persona {n}")
        r = self.client.get("/api/personas/?page_size=2").json()
        self.assertEqual((len(r["results"]), r["count"]), (2, 4))

    def test_usuario_inactivo_pierde_el_acceso(self):
        usuario = self.entrar_como(RolSistema.COORD_RETIRO)
        usuario.activo = False
        usuario.save()
        # Desactivarlo cierra su sesión.
        self.assertEqual(self.client.get(self.ruta).status_code, 401)


class DocumentoYGeografiaTests(Base):
    def crear(self, **datos):
        return self.post("/api/personas/", {
            "primer_nombre": "Ana", "primer_apellido": "Pérez",
            "parroquia": str(self.parroquia.pk), **datos,
        })

    def test_el_documento_se_normaliza_y_no_se_repite(self):
        r = self.crear(tipo_documento="V", documento="v-12.345.678")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(Persona.objects.get().documento, "12345678")

        r = self.crear(tipo_documento="V", documento="12345678")
        self.assertEqual(r.status_code, 400)
        self.assertIn("Ya está registrada con ese documento", r.json()["documento"][0])
        # El mismo número con otro tipo de documento es otra persona.
        self.assertEqual(self.crear(tipo_documento="E", documento="12345678").status_code, 201)

    def test_formato_del_documento(self):
        r = self.crear(tipo_documento="V", documento="12A45")
        self.assertIn("solo números", r.json()["documento"][0])
        r = self.crear(tipo_documento="P", documento="AB1")
        self.assertIn("pasaporte", r.json()["documento"][0])
        self.assertEqual(self.crear(tipo_documento="P", documento="ab123456").status_code, 201)

    def test_edad_minima_de_18(self):
        hoy = timezone.localdate()
        menor = hoy.replace(year=hoy.year - 17).isoformat()
        r = self.crear(fecha_nacimiento=menor)
        self.assertEqual(r.json()["fecha_nacimiento"], ["Debe tener al menos 18 años."])
        adulto = hoy.replace(year=hoy.year - 18, day=1).isoformat()
        self.assertEqual(self.crear(fecha_nacimiento=adulto).status_code, 201)

    def test_sembrar_geografia_no_duplica(self):
        call_command("sembrar_geografia", stdout=open("/dev/null", "w"))
        call_command("sembrar_geografia", stdout=open("/dev/null", "w"))
        venezuela = Pais.objects.get(codigo="VE")
        self.assertEqual(venezuela.estados.count(), 25)
        self.assertTrue(Ciudad.objects.filter(estado__nombre="Miranda", nombre="Los Teques")
                        .exists())
        self.assertEqual(Ciudad.objects.filter(nombre="Caracas").count(), 1)

    def test_la_ciudad_debe_ser_del_estado(self):
        call_command("sembrar_geografia", stdout=open("/dev/null", "w"))
        zulia = Estado.objects.get(nombre="Zulia")
        teques = Ciudad.objects.get(nombre="Los Teques")
        r = self.crear(estado=str(zulia.pk), ciudad=str(teques.pk))
        self.assertEqual(r.json()["ciudad"], ["Esa ciudad no es de ese estado."])

    def test_sembrar_areas_no_duplica(self):
        propia = AreaServicio.objects.create(parroquia=self.parroquia, nombre="Cocina")
        sin_descripcion = AreaServicio.objects.create(nombre="Música / Coro")
        call_command("sembrar_areas", stdout=open("/dev/null", "w"))
        call_command("sembrar_areas", stdout=open("/dev/null", "w"))
        generales = AreaServicio.objects.filter(parroquia=None)
        self.assertEqual(generales.count(), 11)
        self.assertIn("Santísimo", generales.get(nombre="Intercesión / Capilla").descripcion)
        sin_descripcion.refresh_from_db()
        self.assertIn("alabanza", sin_descripcion.descripcion)
        # El área propia de la parroquia no se toca.
        propia.refresh_from_db()
        self.assertEqual(propia.descripcion, "")

    def test_estados_y_ciudades_por_api(self):
        call_command("sembrar_geografia", stdout=open("/dev/null", "w"))
        venezuela = Pais.objects.get(codigo="VE")
        estados = self.client.get(f"/api/estados/?pais={venezuela.pk}&page_size=100").json()
        self.assertEqual(estados["count"], 25)
        miranda = next(e for e in estados["results"] if e["nombre"] == "Miranda")
        ciudades = self.client.get(f"/api/ciudades/?estado={miranda['id']}&page_size=100").json()
        self.assertIn("Los Teques", [c["nombre"] for c in ciudades["results"]])


class NombresYDuplicadosTests(Base):
    def crear(self, **datos):
        return self.post("/api/personas/", {"parroquia": str(self.parroquia.pk), **datos})

    def test_nombre_completo_con_nombres_y_apellidos(self):
        r = self.crear(primer_nombre=" María ", segundo_nombre="José",
                       primer_apellido="Pérez", segundo_apellido="Gómez",
                       nombre="Esto no se usa")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.json()["nombre"], "María José Pérez Gómez")

    def test_pide_primer_nombre_y_primer_apellido(self):
        errores = self.crear(segundo_nombre="José").json()
        self.assertEqual(errores["primer_nombre"], ["Escribe el primer nombre."])
        self.assertEqual(errores["primer_apellido"], ["Escribe el primer apellido."])

    def test_nombres_viejos_se_reparten(self):
        casos = {
            "Ana Pérez": ("Ana", "", "Pérez", ""),
            "José Pérez Gómez": ("José", "", "Pérez", "Gómez"),
            "María José Pérez Gómez": ("María", "José", "Pérez", "Gómez"),
            "Ana María Luisa Pérez Gómez": ("Ana", "María Luisa", "Pérez", "Gómez"),
        }
        for completo, partes in casos.items():
            p = self.persona(completo)
            self.assertEqual((p.primer_nombre, p.segundo_nombre, p.primer_apellido,
                              p.segundo_apellido), partes, completo)
            self.assertEqual(p.nombre, completo)

    def test_no_duplica_por_nombre_y_fecha(self):
        datos = {"primer_nombre": "Ana", "primer_apellido": "Pérez",
                 "fecha_nacimiento": "1990-05-04"}
        self.assertEqual(self.crear(**datos).status_code, 201)
        r = self.crear(**{**datos, "primer_nombre": "ana"})
        self.assertIn("Ya hay una persona con ese nombre y fecha de nacimiento",
                      r.json()["primer_nombre"][0])
        # Aunque la nueva traiga documento, la anterior no tiene: podría ser
        # la misma persona.
        r = self.crear(**datos, documento="12345678")
        self.assertEqual(r.status_code, 400)

    def test_con_documentos_distintos_son_personas_distintas(self):
        datos = {"primer_nombre": "Ana", "primer_apellido": "Pérez",
                 "fecha_nacimiento": "1990-05-04"}
        self.assertEqual(self.crear(**datos, documento="12345678").status_code, 201)
        self.assertEqual(self.crear(**datos, documento="87654321").status_code, 201)

    def test_pasaporte_repetido(self):
        datos = {"primer_nombre": "John", "primer_apellido": "Smith",
                 "tipo_documento": "P", "documento": "AB123456"}
        self.assertEqual(self.crear(**datos).status_code, 201)
        r = self.crear(**{**datos, "primer_nombre": "Juan"})
        self.assertIn("Ya está registrada con ese documento", r.json()["documento"][0])

    def test_ciudad_escrita_si_no_esta_en_la_lista(self):
        colombia = Pais.objects.create(nombre="Colombia", codigo="CO")
        r = self.crear(primer_nombre="Ana", primer_apellido="Gómez", pais=str(colombia.pk),
                       estado_texto="Antioquia", ciudad_texto="Medellín")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(Persona.objects.get().procedencia, "Medellín, Antioquia, Colombia")

    def test_caminante_y_servidor_del_mismo_retiro(self):
        ana, luis = self.persona("Ana Pérez"), self.persona("Luis Gómez")
        Inscripcion.objects.create(retiro=self.retiro, persona=ana)
        r = self.post("/api/participaciones/", {
            "retiro": str(self.retiro.pk), "persona": str(ana.pk), "tipo": "servidor"})
        self.assertEqual(r.json()["persona"], ["Está inscrita como caminante en este retiro."])

        Participacion.objects.create(retiro=self.retiro, persona=luis, tipo="servidor")
        r = self.post("/api/inscripciones/", {
            "retiro": str(self.retiro.pk), "persona": str(luis.pk)})
        self.assertEqual(r.json()["persona"],
                         ["Ya está en el equipo de este retiro como servidor."])
        r = self.post("/api/participaciones/", {
            "retiro": str(self.retiro.pk), "persona": str(luis.pk), "tipo": "coordinador"})
        self.assertEqual(r.json()["persona"], ["Ya está en este retiro como servidor."])

    def test_area_repetida(self):
        luis = Participacion.objects.create(
            retiro=self.retiro, persona=self.persona("Luis Gómez"), tipo="servidor")
        cocina = AreaServicio.objects.create(nombre="Cocina")
        datos = {"participacion": str(luis.pk), "area": str(cocina.pk)}
        self.assertEqual(self.post("/api/servicios/", datos).status_code, 201)
        self.assertEqual(self.post("/api/servicios/", datos).json()["area"],
                         ["Ya sirve en esa área."])
        # Las áreas de un servidor, para editarlo.
        r = self.client.get(f"/api/servicios/?participacion={luis.pk}").json()
        self.assertEqual([x["area_etiqueta"] for x in r["results"]], ["Cocina"])


class PreferenciasTests(Base):
    def cambiar(self, **datos):
        return self.client.patch("/api/auth/preferencias/", datos,
                                 content_type="application/json")

    def test_sin_preferencias_usa_la_moneda_base(self):
        pref = self.client.get("/api/auth/sesion/").json()["preferencias"]
        self.assertEqual(pref, {"retiro": None, "moneda": "USD", "moneda_vista": "USD",
                                "factor": 1, "aviso": "", "tablero_ocultos": []})

    def test_elegir_edicion_y_moneda(self):
        r = self.cambiar(retiro=str(self.retiro.pk), moneda="VES")
        self.assertEqual(r.status_code, 200, r.content)
        pref = r.json()["preferencias"]
        self.assertEqual(pref["retiro"], {"id": str(self.retiro.pk),
                                          "etiqueta": "Emaús hombres #12"})
        # 180 Bs por dólar (el BCV de prueba): un monto en USD se ve x180.
        self.assertEqual((pref["moneda_vista"], pref["factor"]), ("VES", 180.0))
        self.assertEqual(self.cambiar(moneda="EUR").json()["preferencias"]["factor"], 0.9)
        # Cambiar una no borra la otra.
        self.assertEqual(self.client.get("/api/auth/sesion/").json()["preferencias"]
                         ["retiro"]["id"], str(self.retiro.pk))

    def test_columna_de_equivalente_en_la_moneda_elegida(self):
        self.cambiar(moneda="VES")
        (pagos,) = [m for m in self.client.get("/api/esquema/").json() if m["slug"] == "pagos"]
        columna = next(c for c in pagos["lista"] if c["nombre"] == "monto_base")
        self.assertEqual(columna["etiqueta"], "Equivalente en VES")

    def test_sin_tasa_se_ve_en_la_moneda_base(self):
        self.bcv.side_effect = OSError("sin red")
        cache.clear()
        pref = self.cambiar(moneda="VES").json()["preferencias"]
        self.assertEqual((pref["moneda"], pref["moneda_vista"], pref["factor"]),
                         ("VES", "USD", 1))
        self.assertIn("No se pudo consultar la tasa", pref["aviso"])

    def test_valida_lo_que_se_elige(self):
        self.assertEqual(self.cambiar(moneda="MXN").status_code, 400)
        self.assertEqual(
            self.cambiar(retiro="00000000-0000-0000-0000-000000000000").status_code, 400)

    def test_el_tablero_abre_la_edicion_preferida(self):
        otro = Retiro.objects.create(
            parroquia=self.parroquia, tipo=self.retiro.tipo, edicion=13,
            fecha_inicio=date(2027, 1, 10), fecha_fin=date(2027, 1, 12),
        )
        # Sin preferencia: el próximo (el 13, que empieza antes).
        self.assertEqual(self.client.get("/api/resumen/").json()["tablero"]["retiro"]["id"],
                         str(otro.pk))
        self.cambiar(retiro=str(self.retiro.pk))
        self.assertEqual(self.client.get("/api/resumen/").json()["tablero"]["retiro"]["id"],
                         str(self.retiro.pk))

    def test_una_edicion_fuera_del_rol_no_se_puede_elegir(self):
        samuel = Retiro.objects.create(
            parroquia=self.parroquia, tipo=TipoRetiro.objects.create(nombre="Samuel"),
            edicion=1, fecha_inicio=date(2027, 5, 1), fecha_fin=date(2027, 5, 3),
        )
        cuenta = User.objects.create_user("coord", password="x")
        usuario = Usuario.objects.create(persona=self.persona("Coord Uno"), cuenta=cuenta)
        AsignacionRol.objects.create(usuario=usuario, rol=rol_de(RolSistema.COORD_RETIRO), retiro=self.retiro)
        self.client.force_login(cuenta)
        self.assertEqual(self.cambiar(retiro=str(samuel.pk)).status_code, 400)
        self.assertEqual(self.cambiar(retiro=str(self.retiro.pk)).status_code, 200)


class FichaYServidoresTests(Base):
    def setUp(self):
        super().setUp()
        self.ana = self.persona("Ana Pérez", fecha_nacimiento=date(1990, 1, 1))
        self.luis = self.persona("Luis Gómez")
        anterior = Retiro.objects.create(
            parroquia=self.parroquia, tipo=self.retiro.tipo, edicion=11,
            fecha_inicio=date(2026, 3, 6), fecha_fin=date(2026, 3, 8), cuota=50,
            cuota_servidor=20,
        )
        # Ana fue caminante en la 11 (invitada por Luis) y sirve en la 12.
        ins = Inscripcion.objects.create(retiro=anterior, persona=self.ana, padrino=self.luis,
                                         estado=Inscripcion.Estado.CONFIRMADA)
        Pago.objects.create(retiro=anterior, inscripcion=ins, monto=50, concepto="cuota",
                            fecha=date(2026, 3, 1))
        servidora = Participacion.objects.create(retiro=self.retiro, persona=self.ana,
                                                 tipo="servidor")
        cocina = AreaServicio.objects.create(nombre="Cocina")
        musica = AreaServicio.objects.create(nombre="Música")
        AsignacionServicio.objects.create(participacion=servidora, area=cocina,
                                          es_coordinador=True)
        AsignacionServicio.objects.create(participacion=servidora, area=musica)
        # Una participación de caminante no es del equipo.
        Participacion.objects.create(retiro=anterior, persona=self.luis, tipo="caminante")
        Tarea.objects.create(retiro=self.retiro, titulo="Comprar café", responsable=self.ana)

    def test_ficha_junta_todo_lo_de_la_persona(self):
        ficha = self.client.get(f"/api/personas/{self.ana.pk}/ficha/").json()
        self.assertEqual(ficha["persona"]["nombre"], "Ana Pérez")
        self.assertIsNotNone(ficha["persona"]["edad"])
        (caminante,) = ficha["caminante"]
        self.assertEqual((caminante["retiro"]["etiqueta"], caminante["estado"],
                          caminante["padrino"], caminante["estado_pago"]),
                         ("Emaús hombres #11", "confirmada", "Luis Gómez", "pagado"))
        (servidor,) = ficha["servidor"]
        self.assertEqual(servidor["retiro"]["etiqueta"], "Emaús hombres #12")
        self.assertEqual([(a["nombre"], a["es_coordinador"]) for a in servidor["areas"]],
                         [("Cocina", True), ("Música", False)])
        self.assertEqual([t["titulo"] for t in ficha["tareas"]], ["Comprar café"])
        self.assertIsNone(ficha["acceso"])  # no tiene usuario

        luis = self.client.get(f"/api/personas/{self.luis.pk}/ficha/").json()
        self.assertEqual([i["nombre"] for i in luis["invito"]], ["Ana Pérez"])
        self.assertEqual(luis["servidor"], [])  # ser caminante no es servir

    def test_ficha_con_usuario_y_roles(self):
        usuario = Usuario.objects.create(
            persona=self.ana, cuenta=User.objects.create_user("ana", password="x"))
        AsignacionRol.objects.create(usuario=usuario, rol=rol_de(RolSistema.LOGISTICA), retiro=self.retiro)
        acceso = self.client.get(f"/api/personas/{self.ana.pk}/ficha/").json()["acceso"]
        self.assertEqual(acceso["usuario"], "ana")
        self.assertEqual(acceso["roles"], [{"rol": "logistica",
                                            "rol_display": "Coordinador de logística",
                                            "ambito": "Emaús hombres #12"}])

    def test_ficha_respeta_permisos(self):
        user = User.objects.create_user("lector", password="x")
        user.user_permissions.add(Permission.objects.get(codename="view_persona"))
        self.client.force_login(user)
        ficha = self.client.get(f"/api/personas/{self.ana.pk}/ficha/").json()
        self.assertIsNone(ficha["caminante"])
        self.assertIsNone(ficha["servidor"])
        self.assertNotIn("acceso", ficha)
        self.assertNotIn("datos_sensibles", ficha["persona"])
        user.user_permissions.clear()
        self.assertEqual(
            self.client.get(f"/api/personas/{self.ana.pk}/ficha/").status_code, 403)

    def test_servidores_con_sus_areas_y_sin_caminantes(self):
        filas = self.client.get("/api/participaciones/").json()["results"]
        self.assertEqual([(f["persona_etiqueta"], f["areas"]) for f in filas],
                         [("Ana Pérez", "Cocina (coordina), Música")])
        buscadas = self.client.get("/api/participaciones/?search=Música").json()["results"]
        self.assertEqual(len(buscadas), 1)

    def test_menu_por_grupos(self):
        modulos = {m["slug"]: m for m in self.client.get("/api/esquema/").json()}
        self.assertEqual(
            [s for s, m in modulos.items() if m["grupo"] == "Personas"],
            ["personas", "inscripciones", "participaciones", "servicios", "usuarios",
             "asignaciones", "roles"])
        self.assertEqual((modulos["inscripciones"]["titulo"],
                          modulos["participaciones"]["titulo"]), ("Caminantes", "Servidores"))
        self.assertFalse(modulos["servicios"]["en_menu"])
        tipo = next(f for f in modulos["participaciones"]["filtros"] if f["nombre"] == "tipo")
        self.assertEqual([o["valor"] for o in tipo["opciones"]], ["servidor", "coordinador"])


class DisponibilidadServidorTests(Base):
    def registrar(self, nombre, **datos):
        return self.post("/api/participaciones/", {
            "retiro": str(self.retiro.pk), "persona": str(self.persona(nombre).pk),
            "tipo": "servidor", **datos,
        })

    def test_guarda_y_resume_lo_que_indica_el_servidor(self):
        r = self.registrar(
            "Ana Pérez", duerme_en_casa=True,
            disponibilidad={"viernes_noche": True, "sabado_manana": True,
                            "sabado_tarde": True, "sabado_noche": True,
                            "domingo_manana": True},
            camisas={"azul": True, "polo": True}, tallas_camisa={"M": True},
        )
        self.assertEqual(r.status_code, 201, r.content)
        fila = self.client.get("/api/participaciones/").json()["results"][0]
        self.assertEqual(fila["duerme"], "Sí")
        self.assertEqual(fila["turnos"], "Vie: noche · Sáb: todo el día · Dom: mañana")
        self.assertEqual(fila["camisa"], "Azul de servicio, Polo de la Rosa · talla M")

    def test_camisa_sin_talla(self):
        r = self.registrar("Ana Pérez", camisas={"blanca": True})
        self.assertEqual(r.json()["tallas_camisa"], ["Indica la talla de la camisa."])

    def test_resumen_de_logistica_en_el_panel(self):
        self.registrar("Ana Pérez", duerme_en_casa=True, disponibilidad={"viernes_noche": True},
                       camisas={"azul": True}, tallas_camisa={"M": True})
        self.registrar("Luis Gómez", duerme_en_casa=False,
                       disponibilidad={"viernes_noche": True, "domingo_tarde": True},
                       camisas={"azul": True, "polo": True}, tallas_camisa={"XL": True})
        self.registrar("Rosa Díaz")
        log = self.client.get(f"/api/retiros/{self.retiro.pk}/panel/").json()["logistica_equipo"]
        self.assertEqual((log["total"], log["duermen"], log["no_duermen"], log["sin_responder"]),
                         (3, 1, 1, 1))
        turnos = {t["valor"]: t["total"] for t in log["turnos"]}
        self.assertEqual((turnos["viernes_noche"], turnos["domingo_tarde"],
                          turnos["sabado_manana"]), (2, 1, 0))
        self.assertEqual({c["valor"]: c["total"] for c in log["camisas"]},
                         {"azul": 2, "blanca": 0, "polo": 1})
        self.assertEqual({t["valor"]: t["total"] for t in log["tallas"] if t["total"]},
                         {"M": 1, "XL": 1})


class AsistenciaTests(Base):
    def setUp(self):
        super().setUp()
        self.retiro.cuota_servidor = 30
        self.retiro.save()
        cocina = AreaServicio.objects.create(nombre="Cocina")
        self.ana = Participacion.objects.create(
            retiro=self.retiro, persona=self.persona("Ana Pérez"), tipo="servidor",
            duerme_en_casa=True, camisas={"azul": True}, tallas_camisa={"M": True})
        self.luis = Participacion.objects.create(
            retiro=self.retiro, persona=self.persona("Luis Gómez"), tipo="servidor",
            duerme_en_casa=True)
        for p in (self.ana, self.luis):
            AsignacionServicio.objects.create(participacion=p, area=cocina)
        self.panel = f"/api/retiros/{self.retiro.pk}/panel/"

    def cambiar(self, ruta, **datos):
        return self.client.patch(ruta, datos, content_type="application/json")

    def test_un_servidor_empieza_por_confirmar(self):
        self.assertEqual(self.ana.estado, Participacion.Estado.POR_CONFIRMAR)

    def test_confirmar_y_cancelar_un_servidor(self):
        r = self.cambiar(f"/api/participaciones/{self.ana.pk}/", estado="confirmado")
        self.assertEqual(r.json()["estado"], "confirmado")
        r = self.cambiar(f"/api/participaciones/{self.luis.pk}/", estado="cancelado",
                         motivo_cancelacion="Viaje de trabajo")
        self.assertEqual(r.json()["motivo_cancelacion"], "Viaje de trabajo")

        panel = self.client.get(self.panel).json()
        # Quien canceló sigue en la lista (historial), pero no en el equipo activo.
        filas = {f["persona"]: f for f in panel["servidores"]["filas"]}
        self.assertEqual(filas["Luis Gómez"]["estado_display"], "Canceló")
        self.assertEqual(filas["Luis Gómez"]["saldo"], 0)  # ya no debe la cuota
        self.assertEqual(filas["Ana Pérez"]["saldo"], 30)
        self.assertEqual(panel["equipo"]["total"], 1)
        self.assertEqual(panel["equipo"]["areas"][0]["personas"],
                         [{"nombre": "Ana Pérez", "es_coordinador": False}])
        self.assertEqual((panel["logistica_equipo"]["total"],
                          panel["logistica_equipo"]["duermen"]), (1, 1))
        self.assertEqual(panel["servidores"]["por_cobrar"], 30)

        tablero = self.client.get("/api/resumen/").json()["tablero"]["servidores"]
        self.assertEqual((tablero["total"], tablero["confirmados"], tablero["por_confirmar"],
                          tablero["cancelados"]), (1, 1, 0, 1))

        # Al volver, el motivo se borra.
        r = self.cambiar(f"/api/participaciones/{self.luis.pk}/", estado="confirmado")
        self.assertEqual(r.json()["motivo_cancelacion"], "")

    def test_cancelar_un_caminante_con_motivo(self):
        ins = Inscripcion.objects.create(retiro=self.retiro, persona=self.persona("Rosa Díaz"))
        r = self.cambiar(f"/api/inscripciones/{ins.pk}/", estado="cancelada",
                         motivo_cancelacion="Enfermedad")
        self.assertEqual(r.json()["motivo_cancelacion"], "Enfermedad")
        fila = self.client.get(self.panel).json()["caminantes"]["filas"][0]
        self.assertEqual((fila["estado"], fila["motivo_cancelacion"]),
                         ("cancelada", "Enfermedad"))
        r = self.cambiar(f"/api/inscripciones/{ins.pk}/", estado="confirmada")
        self.assertEqual(r.json()["motivo_cancelacion"], "")

    def test_si_cancelo_como_servidor_puede_ser_caminante(self):
        self.cambiar(f"/api/participaciones/{self.ana.pk}/", estado="cancelado")
        r = self.post("/api/inscripciones/", {"retiro": str(self.retiro.pk),
                                              "persona": str(self.ana.persona_id)})
        self.assertEqual(r.status_code, 201, r.content)

    def test_datos_viejos_en_conflicto_no_impiden_cambiar_el_estado(self):
        # Caminante y servidor del mismo retiro, de antes de existir esa regla.
        rosa = self.persona("Rosa Díaz")
        ins = Inscripcion.objects.create(retiro=self.retiro, persona=rosa)
        Participacion.objects.create(retiro=self.retiro, persona=rosa, tipo="servidor")
        r = self.cambiar(f"/api/inscripciones/{ins.pk}/", estado="confirmada")
        self.assertEqual(r.status_code, 200, r.content)
        # El tablero lo avisa, para corregirlo.
        alertas = {a["clave"]: a["total"]
                   for a in self.client.get("/api/resumen/").json()["tablero"]["alertas"]}
        self.assertEqual(alertas["caminante_y_servidor"], 1)
        # Pero no se la puede pasar a otra persona que esté en el equipo.
        r = self.cambiar(f"/api/inscripciones/{ins.pk}/", persona=str(self.ana.persona_id))
        self.assertIn("Ya está en el equipo", r.json()["persona"][0])

    def test_reactivar_si_la_regla_se_vuelve_a_revisar(self):
        # Ana canceló como servidora y se inscribió como caminante: no puede
        # volver al equipo mientras siga inscrita.
        self.cambiar(f"/api/participaciones/{self.ana.pk}/", estado="cancelado")
        Inscripcion.objects.create(retiro=self.retiro, persona=self.ana.persona)
        r = self.cambiar(f"/api/participaciones/{self.ana.pk}/", estado="por_confirmar")
        self.assertEqual(r.json()["persona"], ["Está inscrita como caminante en este retiro."])

    def test_filtrar_servidores_por_estado(self):
        self.cambiar(f"/api/participaciones/{self.ana.pk}/", estado="confirmado")
        r = self.client.get("/api/participaciones/?estado=confirmado").json()
        self.assertEqual([f["persona_etiqueta"] for f in r["results"]], ["Ana Pérez"])


class ReportesTests(Base):
    def test_pdf_exige_sesion_y_permiso(self):
        ruta = f"/retiros/{self.retiro.pk}/lista.pdf"
        self.client.logout()
        r = self.client.get(ruta)
        self.assertEqual(r.status_code, 302)
        self.assertTrue(r["Location"].startswith("/signin"))

        self.client.force_login(User.objects.create_user("otro", password="x"))
        self.assertEqual(self.client.get(ruta).status_code, 403)

    def test_pdf_de_confirmados(self):
        Inscripcion.objects.create(retiro=self.retiro, persona=self.persona("Ana"),
                                   estado=Inscripcion.Estado.CONFIRMADA)
        r = self.client.get(f"/retiros/{self.retiro.pk}/lista.pdf")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertIn("lista-emaus-hombres-12.pdf", r["Content-Disposition"])

        falso = "/retiros/00000000-0000-0000-0000-000000000000/lista.pdf"
        self.assertEqual(self.client.get(falso).status_code, 404)

    def test_importar_personas_no_duplica(self):
        self.persona("Ana Pérez", email="ana@example.com")
        self.persona("Sin correo 1")
        self.persona("Sin correo 2")
        datos = tablib.Dataset(
            ["Ana María Pérez", "", "ana@example.com", "", "San José"],
            ["Luis Gómez", 3001234567, "", "1990-05-01", "San José"],
            ["Otra sin correo", "", "", "", "San José"],
            headers=["nombre", "telefono", "email", "fecha_nacimiento", "parroquia"],
        )
        resultado = PersonaResource().import_data(datos, raise_errors=True)

        self.assertFalse(resultado.has_errors())
        self.assertEqual(Persona.objects.count(), 5)
        self.assertEqual(Persona.objects.get(email="ana@example.com").nombre, "Ana María Pérez")
        luis = Persona.objects.get(nombre="Luis Gómez")
        self.assertEqual((luis.telefono, luis.parroquia), ("3001234567", self.parroquia))

        # Reimportar el mismo archivo no crea a nadie más con correo o teléfono.
        PersonaResource().import_data(datos, raise_errors=True)
        self.assertEqual(Persona.objects.filter(nombre="Luis Gómez").count(), 1)

    def test_importar_casa_por_documento(self):
        self.persona("Ana Pérez", documento="12345678")
        datos = tablib.Dataset(
            ["V", "V-12.345.678", "Ana María Pérez", "San José"],
            headers=["tipo_documento", "documento", "nombre", "parroquia"],
        )
        PersonaResource().import_data(datos, raise_errors=True)
        self.assertEqual(Persona.objects.get().nombre, "Ana María Pérez")

    def test_exportar_personas_omite_datos_sensibles(self):
        self.persona("Ana", datos_sensibles={"direccion": "Calle 1"})
        exportado = PersonaResource().export()
        self.assertEqual(exportado.headers,
                         ["tipo_documento", "documento", "nombre", "primer_nombre",
                          "segundo_nombre", "primer_apellido", "segundo_apellido", "telefono",
                          "email", "fecha_nacimiento", "parroquia"])
        self.assertEqual(exportado[0][10], "San José")


class DemoTests(TestCase):
    def test_sembrar_demo_valida_y_se_rehace_sin_duplicar(self):
        from io import StringIO

        from .models import Inscripcion, Participacion, Persona, Retiro
        call_command("sembrar_demo", stdout=StringIO())
        personas = Persona.objects.count()
        call_command("sembrar_demo", stdout=StringIO())  # ya existe: no hace nada
        self.assertEqual(Persona.objects.count(), personas)
        call_command("sembrar_demo", "--borrar", stdout=StringIO())
        self.assertEqual(Persona.objects.count(), personas)
        retiro = Retiro.objects.get(edicion=12, parroquia__nombre="San José (demo)")
        self.assertEqual(retiro.inscripciones.count(), 24)
        self.assertEqual(retiro.participaciones.count(), 15)
        for modelo in (Persona, Inscripcion, Participacion):
            for objeto in modelo.objects.all():
                objeto.full_clean()
        self.assertTrue(User.objects.get(username="demo.logistica").check_password(
            "demo-betania-123"))


class DescargasTests(Base):
    """Las tablas se descargan en Excel y PDF, con los filtros de la pantalla."""

    def setUp(self):
        super().setUp()
        self.ana = Inscripcion.objects.create(retiro=self.retiro, persona=self.persona("Ana Pérez"),
                                              estado=Inscripcion.Estado.CONFIRMADA)
        Inscripcion.objects.create(retiro=self.retiro, persona=self.persona("Luis Gómez"))
        Pago.objects.create(retiro=self.retiro, inscripcion=self.ana, concepto="cuota",
                            monto=Decimal("18000"), moneda="VES", fecha=date(2027, 3, 1))

    def libro(self, respuesta):
        from io import BytesIO

        from openpyxl import load_workbook
        self.assertEqual(respuesta.status_code, 200, respuesta.content[:300])
        return load_workbook(BytesIO(respuesta.content)).active

    def test_excel_de_un_listado_con_filtros_y_busqueda(self):
        r = self.client.get("/api/inscripciones/exportar/?formato=xlsx&estado=confirmada"
                            "&detalle=Estado: Confirmada")
        self.assertIn(".xlsx", r["Content-Disposition"])
        hoja = self.libro(r)
        self.assertEqual(hoja["A1"].value, "Caminantes")
        self.assertEqual(hoja["A2"].value, "Estado: Confirmada")
        encabezado = [c.value for c in hoja[4]]
        self.assertIn("Estado", encabezado)
        self.assertIn("Celular", encabezado)
        filas = [[c.value for c in fila] for fila in hoja.iter_rows(min_row=5)]
        self.assertEqual(len(filas), 1)
        fila = dict(zip(encabezado, filas[0]))
        self.assertEqual(fila["Estado"], "Confirmada")
        self.assertTrue(str(fila[encabezado[0]]).startswith("Ana Pérez"))

        r = self.client.get("/api/inscripciones/exportar/?formato=xlsx&search=Luis")
        self.assertEqual(self.libro(r).max_row, 5)

    def test_los_montos_quedan_como_numeros(self):
        hoja = self.libro(self.client.get("/api/pagos/exportar/?formato=xlsx"))
        encabezado = [c.value for c in hoja[4]]
        fila = {t: c for t, c in zip(encabezado, hoja[5])}
        self.assertEqual(fila["Monto"].value, 18000)
        self.assertIn("Bs.S", fila["Monto"].number_format)  # en la moneda del pago
        self.assertEqual(fila["Fecha"].number_format, "dd/mm/yyyy")

    def test_pdf_de_un_listado(self):
        r = self.client.get("/api/personas/exportar/?formato=pdf")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertIn("personas-", r["Content-Disposition"])

    def test_exige_permiso_de_ver_y_formato_valido(self):
        self.assertEqual(
            self.client.get("/api/personas/exportar/?formato=doc").status_code, 400)
        self.client.force_login(User.objects.create_user("nadie", password="x"))
        self.assertEqual(
            self.client.get("/api/personas/exportar/?formato=xlsx").status_code, 403)

    def test_el_rol_fija_la_columna_y_la_pone_en_el_subtitulo(self):
        cuenta = User.objects.create_user("coord", password="x")
        usuario = Usuario.objects.create(persona=self.persona("Coord"), cuenta=cuenta)
        AsignacionRol.objects.create(usuario=usuario, rol=rol_de(RolSistema.COORD_RETIRO), retiro=self.retiro)
        self.client.force_login(cuenta)
        hoja = self.libro(self.client.get("/api/retiros/exportar/?formato=xlsx"))
        self.assertNotIn("Parroquia", [c.value for c in hoja[4]])
        self.assertIn("Parroquia: San José", hoja["A2"].value)

    def test_tabla_armada_en_el_navegador(self):
        tabla = {"formato": "xlsx", "titulo": "Servidores · Emaús hombres #12",
                 "columnas": ["Servidor", "Pagado", "Duerme"],
                 "filas": [["Ana", {"monto": 25.5, "moneda": "USD"}, "Sí"], ["Luis", None]]}
        hoja = self.libro(self.post("/api/exportar/", tabla))
        self.assertEqual([c.value for c in hoja[5]], ["Ana", 25.5, "Sí"])
        self.assertEqual(hoja.max_row, 6)
        r = self.post("/api/exportar/", {**tabla, "formato": "pdf"})
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertEqual(self.post("/api/exportar/", {"formato": "xlsx"}).status_code, 400)
        self.client.logout()
        self.assertEqual(self.post("/api/exportar/", tabla).status_code, 401)


class ControlPagosTests(Base):
    """La hoja de control de pagos: monto de cada uno, abonos, pagado y resta."""

    def setUp(self):
        super().setUp()
        self.retiro.cuota_servidor = Decimal("25")
        self.retiro.save()
        self.ana = Inscripcion.objects.create(retiro=self.retiro, persona=self.persona("Ana Pérez"))
        self.luis = Inscripcion.objects.create(retiro=self.retiro, persona=self.persona("Luis Gómez"),
                                               estado=Inscripcion.Estado.CANCELADA)
        self.abel = Participacion.objects.create(retiro=self.retiro, persona=self.persona("Abel Mora"),
                                                 tipo="coordinador", cuota=Decimal("20"))
        self.ramiro = Participacion.objects.create(retiro=self.retiro,
                                                   persona=self.persona("Ramiro Alastre"),
                                                   tipo="servidor")
        AsignacionServicio.objects.create(participacion=self.ramiro,
                                          area=AreaServicio.objects.create(nombre="Intercesión"))
        for monto, dia in [(30, 1), (40, 5)]:
            Pago.objects.create(retiro=self.retiro, inscripcion=self.ana, concepto="cuota",
                                monto=monto, fecha=date(2027, 3, dia))
        Pago.objects.create(retiro=self.retiro, persona=self.ramiro.persona, concepto="aportacion",
                            monto=5, fecha=date(2027, 3, 2))
        self.ruta = f"/api/control-pagos/?retiro={self.retiro.pk}"

    def test_abonos_pagado_y_resta_de_cada_uno(self):
        datos = self.client.get(self.ruta).json()
        ana, luis = datos["caminantes"]["filas"]
        self.assertEqual([a["monto"] for a in ana["abonos"]], [30, 40])  # en orden
        self.assertEqual((ana["cuota"], ana["pagado"], ana["saldo"], ana["estado_pago"]),
                         (100, 70, 30, "parcial"))
        self.assertEqual((luis["saldo"], luis["cancelado"]), (0, True))
        self.assertEqual(datos["caminantes"]["totales"]["esperado"], 100)  # sin el que canceló

        abel, ramiro = datos["servidores"]["filas"]
        # Cada servidor con su monto: 20 el que lo tiene propio, 25 los demás.
        self.assertEqual((abel["cuota"], abel["cuota_propia"], abel["saldo"]), (20, True, 20))
        self.assertEqual((abel["equipo"], ramiro["equipo"]), ("Coordinador general", "Intercesión"))
        self.assertEqual((ramiro["cuota"], ramiro["pagado"], ramiro["saldo"]), (25, 5, 20))
        self.assertEqual(datos["servidores"]["totales"]["saldo"], 40)

    def test_el_panel_usa_el_monto_de_cada_uno(self):
        filas = self.client.get(f"/api/retiros/{self.retiro.pk}/panel/").json()
        abel = next(f for f in filas["servidores"]["filas"] if f["persona"] == "Abel Mora")
        self.assertEqual(abel["saldo"], 20)

    def test_cambiar_el_monto_y_volver_a_la_cuota_del_retiro(self):
        ruta = "/api/control-pagos/monto/"
        r = self.post(ruta, {"tipo": "caminante", "id": str(self.ana.pk), "cuota": 70})
        self.assertEqual(r.json(), {"cuota": 70, "cuota_propia": True})
        ana = self.client.get(self.ruta).json()["caminantes"]["filas"][0]
        self.assertEqual((ana["saldo"], ana["estado_pago"]), (0, "pagado"))  # beca parcial

        r = self.post(ruta, {"tipo": "servidor", "id": str(self.abel.pk), "cuota": None})
        self.assertEqual(r.json(), {"cuota": 25, "cuota_propia": False})
        self.assertEqual(
            self.post(ruta, {"tipo": "servidor", "id": str(self.abel.pk), "cuota": -1}).status_code,
            400)

    def test_solo_quien_ve_y_cobra_los_pagos(self):
        cuenta = User.objects.create_user("logistica", password="x")
        usuario = Usuario.objects.create(persona=self.persona("Logística"), cuenta=cuenta)
        AsignacionRol.objects.create(usuario=usuario, rol=rol_de(RolSistema.LOGISTICA), retiro=self.retiro)
        self.client.force_login(cuenta)
        datos = self.client.get(self.ruta).json()
        self.assertFalse(datos["puede_registrar"])  # logística consulta, no cobra
        self.assertIn("ficha", datos["caminantes"]["filas"][0])
        r = self.post("/api/control-pagos/monto/",
                      {"tipo": "caminante", "id": str(self.ana.pk), "cuota": 10})
        self.assertEqual(r.status_code, 403)

        self.client.force_login(User.objects.create_user("nadie", password="x"))
        self.assertEqual(self.client.get(self.ruta).status_code, 403)


class CuentasEImportacionTests(Base):
    """Cuentas de acceso e importación de personas, sin el admin de Django."""

    def setUp(self):
        super().setUp()
        self.ana = Usuario.objects.create(persona=self.persona("Ana Pérez", email="ana@example.com"))
        self.ruta = f"/api/usuarios/{self.ana.pk}/cuenta/"

    def test_crear_la_cuenta_y_entrar_con_ella(self):
        datos = self.client.get(self.ruta).json()
        self.assertEqual((datos["tiene_cuenta"], datos["email"]), (False, "ana@example.com"))
        r = self.post(self.ruta, {"username": "ana", "password": "corta"})
        self.assertIn("password", r.json())  # la valida Django
        r = self.post(self.ruta, {"username": "ana", "email": "ana@example.com",
                                  "password": "clave-larga-2027"})
        self.assertEqual(r.status_code, 200, r.content)
        self.assertTrue(r.json()["tiene_cuenta"])
        self.client.logout()
        self.assertTrue(self.client.login(username="ana", password="clave-larga-2027"))

    def test_la_contrasena_lleva_al_menos_8_caracteres(self):
        r = self.post(self.ruta, {"username": "ana", "password": "clave27"})
        self.assertIn("8", " ".join(r.json()["password"]))
        r = self.post(self.ruta, {"username": "ana", "password": "clave270"})
        self.assertEqual(r.status_code, 200, r.content)

    def test_la_cedula_no_es_obligatoria(self):
        r = self.post("/api/personas/", {
            "parroquia": str(self.parroquia.pk), "tipo_documento": "V", "documento": "",
            "primer_nombre": "Rosa", "primer_apellido": "Díaz"})
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(Persona.objects.get(primer_nombre="Rosa").documento, "")

    def test_editar_sin_cambiar_la_contrasena_y_desactivar(self):
        self.post(self.ruta, {"username": "ana", "password": "clave-larga-2027"})
        r = self.post(self.ruta, {"username": "ana.perez", "password": ""})
        self.assertEqual(r.json()["username"], "ana.perez")
        cuenta = User.objects.get(username="ana.perez")
        self.assertTrue(cuenta.check_password("clave-larga-2027"))
        # Desactivar al usuario le impide entrar.
        self.client.patch(f"/api/usuarios/{self.ana.pk}/", {"activo": False},
                          content_type="application/json")
        cuenta.refresh_from_db()
        self.assertFalse(cuenta.is_active)
        self.assertIn("ana.perez", str(self.client.get("/api/usuarios/").json()))

    def test_borrar_el_usuario_borra_su_cuenta(self):
        self.post(self.ruta, {"username": "ana", "password": "clave-larga-2027"})
        self.client.delete(f"/api/usuarios/{self.ana.pk}/")
        self.assertFalse(User.objects.filter(username="ana").exists())

    def test_no_repite_nombre_ni_lo_toca_quien_no_puede(self):
        User.objects.create_user("ocupado", password="x")
        r = self.post(self.ruta, {"username": "Ocupado", "password": "clave-larga-2027"})
        self.assertIn("username", r.json())
        self.client.force_login(User.objects.create_user("nadie", password="x"))
        self.assertEqual(self.client.get(self.ruta).status_code, 403)

    def subir(self, contenido, nombre="personas.csv", confirmar=False):
        from django.core.files.uploadedfile import SimpleUploadedFile
        archivo = SimpleUploadedFile(nombre, contenido.encode("utf-8"), "text/csv")
        ruta = "/api/personas/importar/" + ("?confirmar=1" if confirmar else "")
        return self.client.post(ruta, {"archivo": archivo})

    def test_importar_con_vista_previa(self):
        csv = ("tipo_documento,documento,nombre,telefono,parroquia\n"
               "V,12345678,Luis Gómez,0414-1234567,San José\n"
               "V,87654321,Rosa Díaz,,San José\n")
        r = self.subir(csv).json()
        self.assertEqual((r["filas"], r["nuevas"], r["guardado"]), (2, 2, False))
        self.assertFalse(Persona.objects.filter(documento="12345678").exists())  # vista previa
        r = self.subir(csv, confirmar=True).json()
        self.assertTrue(r["guardado"])
        self.assertEqual(Persona.objects.get(documento="12345678").nombre, "Luis Gómez")
        # Volver a subirlo no duplica: las actualiza.
        r = self.subir(csv).json()
        self.assertEqual((r["nuevas"], r["actualizadas"] + r["sin_cambios"]), (0, 2))

    def test_importar_la_plantilla_con_titulo(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        plantilla = self.post("/api/exportar/", {
            "formato": "xlsx", "titulo": "Plantilla de personas",
            "columnas": ["tipo_documento", "documento", "nombre", "parroquia"],
            "filas": [["V", "11222333", "Marta Ruiz", "San José"]],
        }).content
        archivo = SimpleUploadedFile("plantilla.xlsx", plantilla)
        r = self.client.post("/api/personas/importar/?confirmar=1", {"archivo": archivo}).json()
        self.assertEqual((r["filas"], r["nuevas"], r["guardado"]), (1, 1, True), r)
        self.assertEqual(Persona.objects.get(documento="11222333").nombre, "Marta Ruiz")

    def test_importar_sin_parroquia(self):
        csv = "documento,nombre\n12345678,Luis Gómez\n"
        # Con una sola parroquia, se usa esa.
        self.assertEqual(self.subir(csv).json()["nuevas"], 1)
        Parroquia.objects.create(nombre="San Pedro", diocesis=self.parroquia.diocesis)
        self.assertIn("parroquia", self.subir(csv).json()["detail"])
        r = self.subir("documento,nombre,parroquia\n12345678,Luis Gómez,Otra\n").json()
        self.assertEqual(r["errores"], [{"fila": 1, "mensaje": "No existe la parroquia «Otra»."}])

    def test_importar_avisa_si_el_archivo_no_sirve(self):
        self.assertEqual(self.subir("a,b\n1,2\n").status_code, 400)
        self.assertEqual(self.subir("x", nombre="personas.txt").status_code, 400)



class RolesYPermisosTests(Base):
    """Los roles se crean, cambian, asignan y eliminan; sus permisos mandan."""

    def entrar_con(self, rol, **ambito):
        cuenta = User.objects.create_user(f"u{Usuario.objects.count()}", password="x")
        usuario = Usuario.objects.create(persona=self.persona(cuenta.username), cuenta=cuenta)
        AsignacionRol.objects.create(usuario=usuario, rol=rol, **ambito)
        self.client.force_login(cuenta)
        return cuenta

    def permisos(self, rol, lista=None):
        ruta = f"/api/roles/{rol.pk}/permisos/"
        if lista is None:
            return self.client.get(ruta)
        return self.client.put(ruta, {"permisos": lista}, content_type="application/json")

    def test_los_roles_del_sistema_vienen_con_sus_permisos(self):
        logistica = Rol.objects.get(codigo=RolSistema.LOGISTICA)
        self.assertTrue(logistica.sistema and logistica.ve_datos_sensibles)
        codigos = set(logistica.permisos.values_list("codename", flat=True))
        self.assertIn("view_pago", codigos)
        self.assertNotIn("add_pago", codigos)
        self.assertEqual(Rol.objects.filter(sistema=True).count(), 8)

    def test_crear_un_rol_darle_permisos_asignarlo_y_que_mande(self):
        r = self.post("/api/roles/", {"nombre": "Cocina", "ambito": "retiro",
                                      "ve_datos_sensibles": False})
        self.assertEqual(r.status_code, 201, r.content)
        cocina = Rol.objects.get(nombre="Cocina")
        self.assertEqual(cocina.codigo, "cocina")
        tabla = self.permisos(cocina).json()
        self.assertIn("Retiros", [g["grupo"] for g in tabla["grupos"]])
        r = self.permisos(cocina, ["view_retiro", "view_inscripcion", "inventado"])
        self.assertEqual(r.status_code, 200)
        self.assertEqual(set(cocina.permisos.values_list("codename", flat=True)),
                         {"view_retiro", "view_inscripcion"})

        Inscripcion.objects.create(retiro=self.retiro, persona=self.persona("Ana"),
                                   salud_alimentacion={"tratamiento": "Insulina"})
        self.entrar_con(cocina, retiro=self.retiro)
        self.assertEqual(self.client.get("/api/inscripciones/").status_code, 200)
        self.assertEqual(self.client.get("/api/pagos/").status_code, 403)
        fila = self.client.get("/api/inscripciones/").json()["results"][0]
        self.assertNotIn("salud_alimentacion", fila)  # no ve datos confidenciales

    def test_el_ambito_del_rol_se_exige_al_asignarlo(self):
        rol = Rol.objects.create(nombre="Visitante", ambito=Rol.Ambito.RETIRO)
        usuario = Usuario.objects.create(persona=self.persona("Luis"))
        r = self.post("/api/asignaciones/", {"usuario": str(usuario.pk), "rol": str(rol.pk),
                                             "parroquia": str(self.parroquia.pk)})
        self.assertEqual(r.status_code, 400)
        self.assertIn("retiro", r.json())
        r = self.post("/api/asignaciones/", {"usuario": str(usuario.pk), "rol": str(rol.pk),
                                             "retiro": str(self.retiro.pk)})
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(rol.asignaciones.count(), 1)

    def test_no_se_borran_los_del_sistema_ni_los_asignados(self):
        sistema = Rol.objects.get(codigo=RolSistema.TESORERIA)
        r = self.client.delete(f"/api/roles/{sistema.pk}/")
        self.assertEqual(r.status_code, 409)
        self.assertIn("del sistema", r.json()["detail"])
        propio = Rol.objects.create(nombre="Propio", ambito=Rol.Ambito.PARROQUIA)
        usuario = Usuario.objects.create(persona=self.persona("Rosa"))
        AsignacionRol.objects.create(usuario=usuario, rol=propio, parroquia=self.parroquia)
        self.assertIn("asignado", self.client.delete(f"/api/roles/{propio.pk}/").json()["detail"])
        AsignacionRol.objects.all().delete()
        self.assertEqual(self.client.delete(f"/api/roles/{propio.pk}/").status_code, 204)

    def test_nadie_da_permisos_que_no_tiene(self):
        # Quien administra roles sin ser superusuario.
        jefe = Rol.objects.create(nombre="Jefe de roles", ambito=Rol.Ambito.PARROQUIA)
        jefe.permisos.set(Permission.objects.filter(codename__in=[
            "view_rol", "change_rol", "view_retiro"]))
        otro = Rol.objects.create(nombre="Otro", ambito=Rol.Ambito.PARROQUIA)
        self.entrar_con(jefe, parroquia=self.parroquia)
        self.assertEqual(self.permisos(otro, ["view_retiro"]).status_code, 200)
        r = self.permisos(otro, ["view_retiro", "add_pago"])
        self.assertEqual(r.status_code, 403)
        self.assertIn("add_pago", r.json()["detail"])
        # Ni subirse a sí mismo.
        self.assertEqual(self.permisos(jefe, ["view_rol", "change_rol", "view_retiro",
                                              "delete_pago"]).status_code, 403)

    def test_quitar_un_permiso_se_aplica_de_inmediato(self):
        logistica = Rol.objects.get(codigo=RolSistema.LOGISTICA)
        self.permisos(logistica, [c for c in logistica.permisos.values_list("codename", flat=True)
                                  if c != "view_pago"])
        self.entrar_con(logistica, retiro=self.retiro)
        self.assertEqual(self.client.get("/api/pagos/").status_code, 403)


class AuditoriaTests(Base):
    """Quién hizo qué y cuándo; sin los ingresos y salidas del super administrador."""

    def ultimos(self, **filtro):
        from .models import Auditoria
        return list(Auditoria.objects.filter(**filtro).order_by("-fecha"))

    def test_registra_altas_cambios_y_bajas_con_quien_y_que(self):
        from .models import Auditoria
        r = self.post("/api/personas/", {"primer_nombre": "Ana", "primer_apellido": "Pérez",
                                         "parroquia": str(self.parroquia.pk),
                                         "datos_sensibles": {"direccion": "Calle 1"}})
        pk = r.json()["id"]
        alta = self.ultimos(accion="crear", objeto_id=pk)[0]
        self.assertEqual((alta.usuario_id, alta.modelo, alta.parroquia_id),
                         (self.admin.pk, "persona", self.parroquia.pk))
        # Lo confidencial no se copia.
        self.assertEqual(alta.cambios["datos_sensibles"]["confidencial"], True)
        self.assertNotIn("Calle 1", str(alta.cambios))

        self.client.patch(f"/api/personas/{pk}/", {"telefono": "0414-1234567"},
                          content_type="application/json")
        cambio = self.ultimos(accion="cambiar", objeto_id=pk)[0]
        self.assertEqual(cambio.cambios["telefono"]["despues"], "0414-1234567")
        self.assertEqual(set(cambio.cambios), {"telefono"})  # solo lo que cambió

        self.client.delete(f"/api/personas/{pk}/")
        self.assertEqual(self.ultimos(accion="eliminar", objeto_id=pk)[0].objeto, "Ana Pérez")
        self.assertEqual(Auditoria.objects.filter(objeto_id=pk).count(), 3)

    def test_opciones_y_estados_se_guardan_con_su_nombre(self):
        ins = Inscripcion.objects.create(retiro=self.retiro, persona=self.persona("Luis"))
        self.client.patch(f"/api/inscripciones/{ins.pk}/", {"estado": "confirmada"},
                          content_type="application/json")
        cambio = self.ultimos(accion="cambiar", objeto_id=str(ins.pk))[0]
        self.assertEqual((cambio.cambios["estado"]["antes"], cambio.cambios["estado"]["despues"]),
                         ("Solicitada", "Confirmada"))
        fila = self.client.get("/api/auditoria/?accion=cambiar").json()["results"][0]
        self.assertEqual((fila["modulo"], fila["resumen"]),
                         ("Caminantes", "Estado: Solicitada → Confirmada"))

    def test_ingresos_y_salidas_menos_los_del_super_admin(self):
        cuenta = User.objects.create_user("ana", password="clave-larga-2027")
        Usuario.objects.create(persona=self.persona("Ana"), cuenta=cuenta)
        self.client.logout()  # el admin (superusuario): no se registra
        self.assertEqual(self.ultimos(accion="salida"), [])
        self.client.post("/api/auth/login/", {"username": "ana", "password": "mala"},
                         content_type="application/json")
        self.client.post("/api/auth/login/", {"username": "ana", "password": "clave-larga-2027"},
                         content_type="application/json")
        self.client.post("/api/auth/logout/")
        self.client.post("/api/auth/login/", {"username": "admin", "password": "x"},
                         content_type="application/json")
        sesiones = {"accion__in": ["ingreso", "salida", "ingreso_fallido"]}
        acciones = [a.accion for a in self.ultimos(**sesiones)]
        self.assertEqual(acciones, ["salida", "ingreso", "ingreso_fallido"])
        self.assertEqual(self.ultimos(accion="ingreso")[0].usuario_nombre, "Ana (ana)")
        # Tampoco el que tiene un rol global (super administrador).
        jefe = User.objects.create_user("jefe", password="clave-larga-2027")
        AsignacionRol.objects.create(
            usuario=Usuario.objects.create(persona=self.persona("Jefe"), cuenta=jefe),
            rol=rol_de(RolSistema.SUPER_ADMIN))
        self.client.post("/api/auth/login/", {"username": "jefe", "password": "clave-larga-2027"},
                         content_type="application/json")
        self.client.post("/api/auth/logout/")
        self.assertEqual(len(self.ultimos(**sesiones)), 3)

    def test_permisos_de_un_rol_y_filtros(self):
        rol = Rol.objects.create(nombre="Cocina", ambito=Rol.Ambito.RETIRO)
        self.client.put(f"/api/roles/{rol.pk}/permisos/", {"permisos": ["view_retiro"]},
                        content_type="application/json")
        cambio = self.ultimos(accion="cambiar", modelo="rol")[0]
        self.assertEqual(cambio.cambios["permisos"]["detalle"], "agregados")
        hoy = timezone.localdate().isoformat()
        self.assertTrue(self.client.get(f"/api/auditoria/?modelo=rol&desde={hoy}").json()["count"])
        self.assertEqual(self.client.get("/api/auditoria/?hasta=2000-01-01").json()["count"], 0)
        opciones = self.client.get("/api/auditoria/opciones/").json()
        self.assertIn({"valor": "rol", "etiqueta": "Roles y permisos"}, opciones["modulos"])

    def test_no_se_edita_ni_se_borra_y_no_cualquiera_la_ve(self):
        from .models import Auditoria
        self.persona("Rosa")
        registro = Auditoria.objects.first()
        self.assertEqual(self.client.delete(f"/api/auditoria/{registro.pk}/").status_code, 403)
        self.assertEqual(self.client.patch(f"/api/auditoria/{registro.pk}/", {"objeto": "x"},
                                           content_type="application/json").status_code, 403)
        self.client.force_login(User.objects.create_user("nadie", password="x"))
        self.assertEqual(self.client.get("/api/auditoria/").status_code, 403)

    def test_la_demo_no_llena_la_auditoria(self):
        from io import StringIO

        from .models import Auditoria
        antes = Auditoria.objects.count()
        call_command("sembrar_demo", stdout=StringIO())
        self.assertEqual(Auditoria.objects.count(), antes)

    def test_la_ip_no_se_puede_inventar(self):
        from django.test import override_settings
        falsa = {"HTTP_X_FORWARDED_FOR": "6.6.6.6"}
        with override_settings(PROXY_CONFIABLE=False):
            self.client.post("/api/personas/", {"primer_nombre": "Eva", "primer_apellido": "Rey",
                                                "parroquia": str(self.parroquia.pk)},
                             content_type="application/json", **falsa)
        self.assertEqual(self.ultimos(accion="crear", modelo="persona")[0].ip, "127.0.0.1")
        with override_settings(PROXY_CONFIABLE=True):  # detrás de nginx, sí se usa
            self.client.post("/api/personas/", {"primer_nombre": "Iris", "primer_apellido": "Paz",
                                                "parroquia": str(self.parroquia.pk)},
                             content_type="application/json", **falsa)
        self.assertEqual(self.ultimos(accion="crear", modelo="persona")[0].ip, "6.6.6.6")


class FiltrosAvanzadosTests(Base):
    """Panel de filtros: edad, fechas, pago, salud, disponibilidad…"""

    def setUp(self):
        super().setUp()
        hoy = timezone.localdate()
        self.joven = self.persona("Ana Joven", fecha_nacimiento=hoy.replace(year=hoy.year - 25))
        self.mayor = self.persona("Luis Mayor", fecha_nacimiento=hoy.replace(year=hoy.year - 60))
        self.ana = Inscripcion.objects.create(
            retiro=self.retiro, persona=self.joven,
            salud_alimentacion={"tratamiento": "Insulina", "talla_camiseta": "M"})
        self.luis = Inscripcion.objects.create(retiro=self.retiro, persona=self.mayor,
                                               salud_alimentacion={"tratamiento": ""})
        Pago.objects.create(retiro=self.retiro, inscripcion=self.ana, concepto="cuota",
                            monto=100, fecha=date(2027, 3, 1))
        self.retiro.cuota_servidor = Decimal("25")
        self.retiro.save()
        cocina = AreaServicio.objects.create(nombre="Cocina")
        self.rosa = Participacion.objects.create(
            retiro=self.retiro, persona=self.persona("Rosa"), tipo="servidor",
            duerme_en_casa=True, disponibilidad={"viernes_noche": True})
        self.pedro = Participacion.objects.create(
            retiro=self.retiro, persona=self.persona("Pedro"), tipo="servidor")
        AsignacionServicio.objects.create(participacion=self.rosa, area=cocina,
                                          es_coordinador=True)
        Pago.objects.create(retiro=self.retiro, persona=self.rosa.persona, concepto="aportacion",
                            monto=10, fecha=date(2027, 3, 1))
        self.cocina = cocina

    def nombres(self, ruta):
        r = self.client.get(ruta)
        self.assertEqual(r.status_code, 200, r.content)
        return sorted(f["etiqueta"].split(" →")[0].split(" ·")[0] for f in r.json()["results"])

    def test_el_esquema_describe_el_panel(self):
        modulos = {m["slug"]: m for m in self.client.get("/api/esquema/").json()}
        nombres = [f["nombre"] for f in modulos["inscripciones"]["avanzados"]]
        self.assertIn("edad", nombres)
        self.assertIn("tratamiento", nombres)
        self.assertEqual(modulos["retiros"].get("avanzados"), [])

    def test_edad_y_fechas(self):
        self.assertEqual(self.nombres("/api/inscripciones/?f_edad_min=18&f_edad_max=30"),
                         ["Ana Joven"])
        self.assertEqual(self.nombres("/api/personas/?f_edad_min=50"), ["Luis Mayor"])
        hoy = timezone.localdate().isoformat()
        self.assertEqual(len(self.nombres(f"/api/inscripciones/?f_inscrito_desde={hoy}")), 2)
        self.assertEqual(self.nombres("/api/inscripciones/?f_inscrito_hasta=2000-01-01"), [])

    def test_estado_del_pago(self):
        self.assertEqual(self.nombres("/api/inscripciones/?f_estado_pago=pagado"), ["Ana Joven"])
        self.assertEqual(self.nombres("/api/inscripciones/?f_estado_pago=pendiente,parcial"),
                         ["Luis Mayor"])
        self.assertEqual(self.nombres("/api/participaciones/?f_estado_pago=parcial"), ["Rosa"])
        self.assertEqual(self.nombres("/api/participaciones/?f_estado_pago=pendiente"), ["Pedro"])

    def test_salud_solo_para_quien_ve_datos_confidenciales(self):
        self.assertEqual(self.nombres("/api/inscripciones/?f_tratamiento=si"), ["Ana Joven"])
        self.assertEqual(self.nombres("/api/inscripciones/?f_tratamiento=no"), ["Luis Mayor"])
        self.assertEqual(self.nombres("/api/inscripciones/?f_talla=M"), ["Ana Joven"])
        cuenta = User.objects.create_user("teso", password="x")
        u = Usuario.objects.create(persona=self.persona("Teso"), cuenta=cuenta)
        AsignacionRol.objects.create(usuario=u, rol=rol_de(RolSistema.TESORERIA),
                                     parroquia=self.parroquia)
        self.client.force_login(cuenta)
        # Sin acceso a datos confidenciales, el filtro no existe (se ignora).
        self.assertEqual(len(self.nombres("/api/inscripciones/?f_tratamiento=si")), 2)
        modulos = {m["slug"]: m for m in self.client.get("/api/esquema/").json()}
        self.assertNotIn("tratamiento", [f["nombre"] for f in modulos["inscripciones"]["avanzados"]])

    def test_servicio_y_disponibilidad(self):
        self.assertEqual(self.nombres(f"/api/participaciones/?f_area={self.cocina.pk}"), ["Rosa"])
        self.assertEqual(self.nombres("/api/participaciones/?f_coordina=si"), ["Rosa"])
        self.assertEqual(self.nombres("/api/participaciones/?f_sin_area=si"), ["Pedro"])
        self.assertEqual(self.nombres("/api/participaciones/?f_duerme=sin"), ["Pedro"])
        self.assertEqual(self.nombres("/api/participaciones/?f_turno=viernes_noche,sabado_manana"),
                         ["Rosa"])

    def test_participacion_de_las_personas(self):
        self.assertEqual(self.nombres("/api/personas/?f_caminante=si"), ["Ana Joven", "Luis Mayor"])
        self.assertEqual(self.nombres("/api/personas/?f_servidor=si"), ["Pedro", "Rosa"])
        self.assertEqual(len(self.nombres(f"/api/personas/?f_en_retiro={self.retiro.pk}")), 4)

    def test_valores_raros_y_descarga(self):
        self.assertEqual(self.nombres("/api/inscripciones/?f_edad_min=abc"), [])
        self.assertEqual(self.nombres("/api/participaciones/?f_area=no-es-un-id"), [])
        self.assertEqual(len(self.nombres("/api/inscripciones/?f_estado_pago=inventado")), 2)
        r = self.client.get("/api/inscripciones/exportar/?formato=xlsx&f_edad_max=30")
        from io import BytesIO

        from openpyxl import load_workbook
        hoja = load_workbook(BytesIO(r.content)).active
        self.assertEqual(hoja.max_row, 5)  # título, subtítulo, vacía, encabezado y Ana


class PaginasDeErrorTests(Base):
    def test_404_y_500_en_espanol(self):
        from django.test import override_settings
        with override_settings(DEBUG=False):
            r = self.client.get("/retiros/00000000-0000-0000-0000-000000000000/lista.pdf")
        self.assertEqual(r.status_code, 404)
        self.assertIn("No encontramos esta página", r.content.decode())
        from django.template.loader import render_to_string
        self.assertIn("Sistema no disponible", render_to_string("500.html"))


def _imagen(lineas, ancho=2000, columnas=None, cuadricula=False):
    """Una planilla o un listado de prueba, como PNG. `cuadricula`: con las
    líneas de las celdas, como una hoja de Excel."""
    import io

    from PIL import Image, ImageDraw, ImageFont
    fuente = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 32)
    img = Image.new("RGB", (ancho, 140 + 70 * len(lineas)), "white")
    d = ImageDraw.Draw(img)
    if cuadricula:
        abajo = 45 + 70 * len(lineas)
        for i in range(len(lineas) + 1):
            d.line((20, 45 + 70 * i, ancho - 20, 45 + 70 * i), fill="black", width=3)
        for x in [20, *(c - 15 for c in columnas[1:]), ancho - 20]:
            d.line((x, 45, x, abajo), fill="black", width=3)
    for i, linea in enumerate(lineas):
        if columnas:
            for x, celda in zip(columnas, linea):
                d.text((x, 60 + 70 * i), celda, fill="black", font=fuente)
        else:
            d.text((80, 60 + 70 * i), linea, fill="black", font=fuente)
    salida = io.BytesIO()
    img.save(salida, "PNG")
    return salida.getvalue()


PLANILLA = [
    "PLANILLA DE INSCRIPCIÓN",
    "Nombres y Apellidos: JOSÉ GREGORIO PÉREZ LOZADA",
    "Cédula de Identidad: V-12.345.678     Fecha de Nacimiento: 12/03/1980",
    "Teléfono Celular: 0414-123.45.67     Teléfono de Habitación: 0212 555 1234",
    "Estado: Miranda     Ciudad: Los Teques",
    "Estado Civil: Casado     Número de hijos: 3",
    "Tratamiento: Metformina 850 mg",
    "Talla de camisa: L",
    "En caso de emergencia avisar a: María Rodríguez",
    "Parentesco: Esposa     Celular: 0424-765.43.21",
    "Invitado por: Pedro Ramírez",
]


class OCRTests(Base):
    """Leer planillas y listados con Tesseract (en el propio servidor)."""

    def subir(self, ruta, contenido, nombre="planilla.png", **datos):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return self.client.post(ruta, {"archivo": SimpleUploadedFile(nombre, contenido), **datos})

    def test_planilla_de_caminante(self):
        call_command("sembrar_geografia", verbosity=0)
        r = self.subir("/api/ocr/planilla/", _imagen(PLANILLA), tipo="caminante")
        self.assertEqual(r.status_code, 200, r.content)
        c = {k: v["valor"] for k, v in r.json()["campos"].items()}
        self.assertEqual((c["primer_nombre"], c["segundo_nombre"], c["primer_apellido"],
                          c["segundo_apellido"]), ("José", "Gregorio", "Pérez", "Lozada"))
        self.assertEqual((c["tipo_documento"], c["documento"]), ("V", "12345678"))
        self.assertEqual(c["fecha_nacimiento"], "1980-03-12")
        self.assertEqual((c["telefono"], c["telefono_casa"]), ("0414-1234567", "0212-5551234"))
        self.assertEqual((c["estado_civil"], c["hijos"], c["talla_camiseta"]), ("Casado", "3", "L"))
        self.assertEqual(c["tratamiento"], "Metformina 850 mg")
        # Los teléfonos después de "En caso de emergencia" son del familiar.
        self.assertEqual((c["contacto_nombre"], c["contacto_relacion"], c["contacto_celular"]),
                         ("María Rodríguez", "Esposa", "0424-7654321"))
        self.assertEqual(c["padrino_nombre"], "Pedro Ramírez")
        # Estado y ciudad, reconocidos en el catálogo.
        self.assertEqual(Ciudad.objects.get(pk=c["ciudad"]).nombre, "Los Teques")
        self.assertTrue(all(0 <= v["confianza"] <= 100 for v in r.json()["campos"].values()))

    def test_planilla_de_servidor_y_pdf(self):
        import pymupdf
        call_command("sembrar_areas", verbosity=0)
        imagen = _imagen(["PLANILLA DE SERVIDOR", "Nombres: Rosa María   Apellidos: Díaz Peña",
                          "Área de servicio: Cocina",
                          "Dormirá en la casa: Sí", "Viernes: noche   Sábado: todo el día",
                          "Camisa: Azul   Talla: M"])
        documento = pymupdf.open()
        pagina = documento.new_page()
        pagina.insert_image(pagina.rect, stream=imagen)
        r = self.subir("/api/ocr/planilla/", documento.tobytes(), "planilla.pdf", tipo="servidor")
        self.assertEqual(r.status_code, 200, r.content)
        c = {k: v["valor"] for k, v in r.json()["campos"].items()}
        self.assertEqual((c["primer_nombre"], c["primer_apellido"]), ("Rosa", "Díaz"))
        self.assertEqual(AreaServicio.objects.get(pk=c["areas"][0]).nombre, "Cocina")
        self.assertEqual(c["duerme"], "si")
        self.assertIn("viernes_noche", c["turnos"])
        self.assertIn("sabado_tarde", c["turnos"])
        self.assertEqual((c["camisas"], c["tallas"]), (["azul"], ["M"]))

    def test_sin_datos_confidenciales_no_los_devuelve(self):
        cuenta = User.objects.create_user("teso", password="x")
        u = Usuario.objects.create(persona=self.persona("Teso"), cuenta=cuenta)
        rol = Rol.objects.create(nombre="Inscribe", ambito=Rol.Ambito.PARROQUIA)
        rol.permisos.set(Permission.objects.filter(codename__in=["add_persona", "view_persona"]))
        AsignacionRol.objects.create(usuario=u, rol=rol, parroquia=self.parroquia)
        self.client.force_login(cuenta)
        campos = self.subir("/api/ocr/planilla/", _imagen(PLANILLA)).json()["campos"]
        self.assertIn("documento", campos)
        for clave in ("tratamiento", "contacto_nombre", "estado_civil", "telefono_casa"):
            self.assertNotIn(clave, campos)

    def test_archivos_que_no_sirven_y_permisos(self):
        self.assertEqual(self.subir("/api/ocr/planilla/", b"hola", "nota.txt").status_code, 400)
        self.assertEqual(self.subir("/api/ocr/planilla/", b"no es imagen").status_code, 400)
        self.client.force_login(User.objects.create_user("nadie", password="x"))
        self.assertEqual(self.subir("/api/ocr/planilla/", _imagen(PLANILLA)).status_code, 403)

    LISTADO = [
        ["N°", "NOMBRE Y APELLIDO", "CELULAR", "EDAD", "TALLA", "CONTACTO", "CELULAR"],
        ["1", "ABEL MORA", "0414-1234567", "45", "L", "ROSA MORA", "0424-7654321"],
        ["2", "LUIS ALASTRE", "0412-2223344", "38", "M", "ANA PEÑA", "0416-9998877"],
    ]
    COLUMNAS = [60, 150, 680, 960, 1080, 1200, 1620]

    def test_listado_y_carga_sin_duplicar(self):
        r = self.subir("/api/ocr/listado/", _imagen(self.LISTADO, columnas=self.COLUMNAS),
                       "lista.png")
        self.assertEqual(r.status_code, 200, r.content)
        filas = r.json()["filas"]
        self.assertEqual([f["nombre"] for f in filas], ["Abel Mora", "Luis Alastre"])
        self.assertEqual((filas[0]["telefono"], filas[0]["edad"], filas[0]["talla"],
                          filas[0]["contacto_nombre"], filas[0]["contacto_celular"]),
                         ("0414-1234567", 45, "L", "Rosa Mora", "0424-7654321"))

        datos = {"retiro": str(self.retiro.pk), "como": "caminantes",
                 "filas": filas + [{"nombre": "Solo"}]}
        r = self.client.post("/api/ocr/listado/importar/", datos, content_type="application/json")
        self.assertEqual((r.json()["creados"], r.json()["errores"]), (2, 1), r.content)
        abel = Inscripcion.objects.get(persona__primer_nombre="Abel")
        self.assertEqual((abel.salud_alimentacion["talla_camiseta"],
                          abel.contacto_emergencia["celular"]), ("L", "0424-7654321"))
        # La misma hoja otra vez: no duplica personas ni inscripciones.
        r = self.client.post("/api/ocr/listado/importar/", datos, content_type="application/json")
        self.assertEqual((r.json()["creados"], r.json()["existentes"]), (0, 2))
        self.assertEqual(Persona.objects.filter(primer_nombre="Abel").count(), 1)

    CONTROL = [
        ["N°", "NOMBRE SERVIDOR", "EQUIPO", "MONTO $", "ABONO $", "ABONO $", "RESTA", "TOTAL"],
        ["1", "ABEL MORA", "COCINA", "20", "", "", "20", "20"],
        ["2", "RAMIRO ALASTRE", "MUSICA", "25", "5", "", "20", "25"],
    ]
    COLUMNAS_CONTROL = [50, 130, 640, 930, 1140, 1350, 1560, 1770]

    def test_control_de_pago_con_cuadricula(self):
        imagen = _imagen(self.CONTROL, ancho=1990, columnas=self.COLUMNAS_CONTROL,
                         cuadricula=True)
        r = self.subir("/api/ocr/listado/", imagen, "control.png")
        self.assertEqual(r.status_code, 200, r.content)
        filas = r.json()["filas"]
        self.assertEqual([(f["nombre"], f.get("monto"), f.get("abono")) for f in filas],
                         [("Abel Mora", "20", None), ("Ramiro Alastre", "25", "5")])
        self.assertIn("monto", r.json()["columnas"])

        datos = {"retiro": str(self.retiro.pk), "como": "servidores", "filas": filas}
        r = self.client.post("/api/ocr/listado/importar/", datos, content_type="application/json")
        self.assertEqual(r.json()["creados"], 2, r.content)
        ramiro = Participacion.objects.get(persona__primer_nombre="Ramiro")
        self.assertEqual(ramiro.cuota, Decimal("25"))
        pago = Pago.objects.get(persona=ramiro.persona)
        self.assertEqual((pago.monto, pago.concepto), (Decimal("5"), Pago.Concepto.APORTACION))
        self.assertFalse(Pago.objects.filter(persona__primer_nombre="Abel").exists())

        # La misma hoja otra vez, con un abono más: solo se registra lo nuevo.
        filas[1]["abono"] = "15"
        r = self.client.post("/api/ocr/listado/importar/", datos, content_type="application/json")
        self.assertEqual(r.json()["existentes"], 2, r.content)
        self.assertEqual(sorted(Pago.objects.filter(persona=ramiro.persona)
                                .values_list("monto", flat=True)),
                         [Decimal("5"), Decimal("10")])
        r = self.client.post("/api/ocr/listado/importar/", datos, content_type="application/json")
        self.assertEqual(Pago.objects.filter(persona=ramiro.persona).count(), 2)
        self.assertIn("ya estaba registrado", str(r.json()["resultados"]))

    def test_listado_ilegible(self):
        r = self.subir("/api/ocr/listado/", _imagen(["una hoja sin tabla"]), "x.png")
        self.assertEqual(r.status_code, 400)
        self.assertIn("encabezados", r.json()["detail"])


class EnlacesDelTableroTests(Base):
    """Cada número del tablero coincide con la lista a la que lleva."""

    def test_los_filtros_del_tablero_dan_el_mismo_numero(self):
        r = self.retiro
        Inscripcion.objects.create(retiro=r, persona=self.persona("A"),
                                   salud_alimentacion={"tratamiento": "Insulina"})
        Inscripcion.objects.create(retiro=r, persona=self.persona("B"),
                                   salud_alimentacion={"condicion_medica": "Asma",
                                                       "regimen_alimenticio": "Sin sal"})
        Inscripcion.objects.create(retiro=r, persona=self.persona("C"),
                                   estado=Inscripcion.Estado.CANCELADA,
                                   salud_alimentacion={"tratamiento": "X"})
        Participacion.objects.create(retiro=r, persona=self.persona("S1"), tipo="servidor")
        Participacion.objects.create(retiro=r, persona=self.persona("S2"), tipo="servidor",
                                     estado=Participacion.Estado.CANCELADO)
        tablero = self.client.get(f"/api/resumen/?retiro={r.pk}").json()["tablero"]
        atencion = tablero["caminantes"]["atencion"]
        contar = lambda ruta: self.client.get(ruta).json()["count"]  # noqa: E731
        base = f"/api/inscripciones/?retiro={r.pk}&f_vigente=si"
        self.assertEqual(contar(f"{base}&f_salud=si"), atencion["salud"])
        self.assertEqual(atencion["salud"], 2)
        self.assertEqual(contar(f"{base}&f_comida=si"), atencion["comida"])
        self.assertEqual(contar(f"/api/participaciones/?retiro={r.pk}&f_vigente=si"),
                         tablero["servidores"]["total"])


class AsignacionCoherenteTests(Base):
    def test_el_retiro_manda_sobre_la_parroquia(self):
        otra = Parroquia.objects.create(nombre="Otra", diocesis=self.parroquia.diocesis)
        cuenta = User.objects.create_user("coord", password="x")
        usuario = Usuario.objects.create(persona=self.persona("Coord"), cuenta=cuenta)
        rol = rol_de(RolSistema.COORD_RETIRO)
        # Ya no se puede guardar un retiro con la parroquia de otro lado…
        r = self.post("/api/asignaciones/", {"usuario": str(usuario.pk), "rol": str(rol.pk),
                                             "retiro": str(self.retiro.pk),
                                             "parroquia": str(otra.pk)})
        self.assertEqual(r.status_code, 400)
        self.assertIn("parroquia", r.json())
        # …y si ya existía así, se usa la parroquia del retiro.
        AsignacionRol.objects.create(usuario=usuario, rol=rol, retiro=self.retiro, parroquia=otra)
        self.client.force_login(cuenta)
        self.assertEqual(self.client.get("/api/retiros/").json()["count"], 1)


class MiCuentaTests(Base):
    def setUp(self):
        super().setUp()
        self.cuenta = User.objects.create_user("ana", password="clave-larga-2027")
        usuario = Usuario.objects.create(persona=self.persona("Ana Pérez"), cuenta=self.cuenta)
        AsignacionRol.objects.create(usuario=usuario, rol=rol_de(RolSistema.LOGISTICA),
                                     retiro=self.retiro)
        self.client.force_login(self.cuenta)

    def test_la_sesion_trae_su_nombre_ficha_y_roles(self):
        s = self.client.get("/api/auth/sesion/").json()
        self.assertEqual(s["nombre"], "Ana Pérez")
        self.assertEqual(s["persona"]["nombre"], "Ana Pérez")
        self.assertEqual(s["roles"], ["Coordinador de logística · Emaús hombres #12"])

    def test_cambiar_correo(self):
        r = self.client.patch("/api/auth/cuenta/", {"email": "no es correo"},
                              content_type="application/json")
        self.assertEqual(r.status_code, 400)
        r = self.client.patch("/api/auth/cuenta/", {"email": "ana@example.com"},
                              content_type="application/json")
        self.assertEqual(r.json()["email"], "ana@example.com")

    def test_cambiar_contrasena_sin_salir(self):
        ruta = "/api/auth/cuenta/clave/"
        r = self.post(ruta, {"actual": "mala", "nueva": "otra-clave-2027"})
        self.assertIn("actual", r.json())
        r = self.post(ruta, {"actual": "clave-larga-2027", "nueva": "123"})
        self.assertIn("nueva", r.json())
        r = self.post(ruta, {"actual": "clave-larga-2027", "nueva": "otra-clave-larga-2027"})
        self.assertEqual(r.status_code, 200)
        # Sigue con sesión y la nueva sirve.
        self.assertTrue(self.client.get("/api/auth/sesion/").json()["autenticado"])
        self.client.logout()
        self.assertTrue(self.client.login(username="ana", password="otra-clave-larga-2027"))

    def test_mis_tareas(self):
        persona = self.cuenta.usuario_retiros.persona
        Tarea.objects.create(retiro=self.retiro, titulo="Mía", responsable=persona)
        Tarea.objects.create(retiro=self.retiro, titulo="Otra")
        r = self.client.get(f"/api/tareas/?responsable={persona.pk}").json()
        self.assertEqual([t["titulo"] for t in r["results"]], ["Mía"])


class SesionVencidaTests(Base):
    def test_sin_sesion_es_401_y_sin_permiso_403(self):
        cuenta = User.objects.create_user("ana", password="clave-larga-2027")
        self.client.force_login(cuenta)
        self.assertEqual(self.client.get("/api/pagos/").status_code, 403)  # sin rol
        # Cambiar la contraseña en otra sesión cierra esta.
        cuenta.set_password("otra-clave-larga-2027")
        cuenta.save()
        self.assertEqual(self.client.get("/api/pagos/").status_code, 401)
        self.assertEqual(self.client.get("/api/auth/cuenta/").status_code, 401)


class LimiteDeIntentosTests(Base):
    """Bloqueo por intentos fallidos y otros refuerzos de seguridad."""

    def setUp(self):
        super().setUp()
        self.client.logout()
        self.ana = User.objects.create_user("ana", password="clave-larga-2027")
        Usuario.objects.create(persona=self.persona("Ana"), cuenta=self.ana)

    def entrar(self, usuario="ana", clave="clave-larga-2027", ip="10.0.0.1"):
        return self.client.post("/api/auth/login/", {"username": usuario, "password": clave},
                                content_type="application/json", REMOTE_ADDR=ip)

    def test_cinco_fallos_bloquean_aunque_luego_acierte(self):
        from .models import Auditoria
        for _ in range(3):
            r = self.entrar(clave="mala")
        self.assertIn("Te quedan 2 intentos", r.json()["detail"])
        self.entrar(clave="mala")
        r = self.entrar(clave="mala")
        self.assertEqual(r.status_code, 429)
        r = self.entrar()  # la correcta: igual bloqueado
        self.assertEqual(r.status_code, 429)
        self.assertIn("15 minutos", r.json()["detail"])
        self.assertTrue(Auditoria.objects.filter(accion="bloqueo", usuario_nombre="ana").exists())
        # Otra cuenta desde otra IP sí puede entrar.
        User.objects.create_user("luis", password="clave-larga-2027")
        self.assertEqual(self.entrar("luis", ip="10.0.0.2").status_code, 200)

    def test_un_usuario_que_no_existe_se_bloquea_igual(self):
        for _ in range(5):
            r = self.entrar("no.existe", "x")
        self.assertEqual(r.status_code, 429)  # no revela si existe

    def test_limite_por_ip(self):
        for n in range(20):
            self.entrar(f"usuario{n}", "x", ip="10.9.9.9")
        self.assertEqual(self.entrar(ip="10.9.9.9").status_code, 429)
        self.assertEqual(self.entrar(ip="10.0.0.3").status_code, 200)

    def test_el_bloqueo_vence_y_entrar_limpia_los_fallos(self):
        from datetime import timedelta

        from .models import IntentoFallido
        for _ in range(5):
            self.entrar(clave="mala")
        IntentoFallido.objects.update(fecha=timezone.now() - timedelta(minutes=16))
        self.assertEqual(self.entrar().status_code, 200)
        self.assertFalse(IntentoFallido.objects.filter(usuario="ana").exists())

    def test_tambien_en_el_admin_de_django(self):
        admin = User.objects.create_superuser("jefe", "j@example.com", "clave-larga-2027")
        for _ in range(5):
            self.client.post("/admin/login/", {"username": "jefe", "password": "mala"},
                             REMOTE_ADDR="10.0.0.5")
        self.client.post("/admin/login/", {"username": "jefe", "password": "clave-larga-2027"},
                         REMOTE_ADDR="10.0.0.5")
        self.assertNotIn("_auth_user_id", self.client.session)
        self.assertTrue(admin.check_password("clave-larga-2027"))

    def test_desbloquear_desde_usuarios(self):
        for _ in range(5):
            self.entrar(clave="mala")
        self.client.force_login(self.admin)
        usuario = self.ana.usuario_retiros
        cuenta = self.client.get(f"/api/usuarios/{usuario.pk}/cuenta/").json()
        self.assertGreater(cuenta["bloqueado_minutos"], 0)
        r = self.post(f"/api/usuarios/{usuario.pk}/cuenta/desbloquear/", {})
        self.assertEqual(r.json()["bloqueado_minutos"], 0)
        self.client.logout()
        self.assertEqual(self.entrar().status_code, 200)

    def test_contrasenas_de_al_menos_8_y_cabeceras(self):
        self.client.force_login(self.ana)
        r = self.post("/api/auth/cuenta/clave/", {"actual": "clave-larga-2027",
                                                  "nueva": "Corta-1"})
        self.assertIn("nueva", r.json())
        r = self.client.get("/api/auth/sesion/")
        self.assertEqual(r["X-Frame-Options"], "DENY")
        self.assertEqual(r["X-Content-Type-Options"], "nosniff")
        self.assertEqual(r["Referrer-Policy"], "same-origin")


class CamaYMesaTests(Base):
    """Número de cama y de mesa del caminante: opcionales, por retiro."""

    def test_opcionales_se_guardan_y_salen_en_la_lista(self):
        persona = self.persona("Abel Mora")
        r = self.post("/api/inscripciones/", {"retiro": str(self.retiro.pk),
                                              "persona": str(persona.pk)})
        self.assertEqual(r.status_code, 201, r.content)  # sin cama ni mesa
        ins = Inscripcion.objects.get(persona=persona)
        self.assertEqual((ins.cama, ins.mesa), ("", ""))
        r = self.client.patch(f"/api/inscripciones/{ins.pk}/", {"cama": "12", "mesa": "B-3"},
                              content_type="application/json")
        self.assertEqual(r.status_code, 200, r.content)
        fila = self.client.get("/api/inscripciones/").json()["results"][0]
        self.assertEqual((fila["cama"], fila["mesa"]), ("12", "B-3"))
        r = self.client.patch(f"/api/inscripciones/{ins.pk}/", {"cama": "x" * 11},
                              content_type="application/json")
        self.assertEqual(r.status_code, 400)


class InactividadTests(Base):
    """La sesión se cierra tras 2 horas sin actividad del usuario."""

    def hace(self, minutos):
        from .seguridad import ACTIVIDAD
        sesion = self.client.session
        sesion[ACTIVIDAD] = int(timezone.now().timestamp()) - minutos * 60
        sesion.save()

    def actividad(self):
        from .seguridad import ACTIVIDAD
        return self.client.session.get(ACTIVIDAD)

    def test_la_sesion_trae_el_limite(self):
        self.assertEqual(self.client.get("/api/auth/sesion/").json()["inactividad_minutos"], 120)

    def test_dos_horas_sin_actividad_cierran_la_sesion(self):
        self.hace(100)
        self.assertEqual(self.client.get("/api/retiros/").status_code, 200)  # aún no
        self.hace(121)
        self.assertEqual(self.client.get("/api/retiros/").status_code, 401)
        self.assertFalse(self.client.get("/api/auth/sesion/").json()["autenticado"])

    def test_lo_que_la_app_consulta_sola_no_cuenta(self):
        self.hace(100)
        antes = self.actividad()
        r = self.client.get("/api/resumen/", HTTP_X_SEGUNDO_PLANO="1")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self.actividad(), antes)  # no se renovó
        self.hace(121)
        r = self.client.get("/api/resumen/", HTTP_X_SEGUNDO_PLANO="1")
        self.assertEqual(r.status_code, 401)

    def test_seguir_conectado_renueva(self):
        self.hace(119)
        r = self.client.post("/api/auth/actividad/")
        self.assertEqual(r.status_code, 200)
        self.assertGreater(self.actividad(), int(timezone.now().timestamp()) - 60)
        self.assertEqual(self.client.get("/api/retiros/").status_code, 200)


class SaludTests(TestCase):
    """/api/salud/: lo consulta el cron que mantiene vivo el servidor."""

    def test_responde_sin_sesion(self):
        r = self.client.get("/api/salud/")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json(), {"estado": "ok"})
        self.assertEqual(r["Cache-Control"], "no-store")
        self.assertNotIn("sessionid", r.cookies)

    def test_avisa_si_la_base_no_responde(self):
        with mock.patch("retiros.api.salud._consultar_base", side_effect=Exception):
            r = self.client.get("/api/salud/")
        self.assertEqual(r.status_code, 503)

    def test_solo_get(self):
        self.assertEqual(self.client.post("/api/salud/").status_code, 405)


class FotoPersonaTests(Base):
    """Foto de perfil: se guarda recortada, sale en los listados y la sesión."""

    def setUp(self):
        super().setUp()
        self.ana = self.persona("Ana Pérez", sexo="F")

    def archivo(self, ancho=900, alto=600, formato="PNG", modo="RGBA"):
        import io

        from django.core.files.uploadedfile import SimpleUploadedFile
        from PIL import Image
        datos = io.BytesIO()
        Image.new(modo, (ancho, alto), (200, 30, 30, 128) if modo == "RGBA" else "red").save(
            datos, formato)
        return SimpleUploadedFile(f"foto.{formato.lower()}", datos.getvalue())

    def subir(self, archivo):
        return self.client.post(f"/api/personas/{self.ana.pk}/foto/", {"foto": archivo})

    def test_subir_la_recorta_y_la_sirve(self):
        import io

        from PIL import Image
        r = self.subir(self.archivo())
        self.assertEqual(r.status_code, 200, r.content)
        url = r.json()["foto"]
        self.assertTrue(url.startswith(f"/api/personas/{self.ana.pk}/foto/?v="))
        r = self.client.get(url)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], "image/jpeg")
        self.assertIn("immutable", r["Cache-Control"])
        imagen = Image.open(io.BytesIO(r.content))
        self.assertEqual((imagen.format, imagen.size), ("JPEG", (400, 400)))

    def test_no_acepta_lo_que_no_es_imagen(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        r = self.subir(SimpleUploadedFile("foto.jpg", b"no soy una foto"))
        self.assertEqual(r.status_code, 400)
        self.assertIn("No es una imagen válida", r.json()["foto"][0])
        self.ana.refresh_from_db()
        self.assertIsNone(self.ana.foto_actualizada)

    def test_sale_en_listados_y_en_la_sesion(self):
        from .models import Inscripcion, Usuario
        self.subir(self.archivo(formato="JPEG", modo="RGB"))
        self.ana.refresh_from_db()
        fila = self.client.get("/api/personas/").json()["results"][0]
        self.assertEqual((fila["foto"], fila["sexo"]), (self.ana.foto_url, "F"))
        Inscripcion.objects.create(retiro=self.retiro, persona=self.ana)
        fila = self.client.get("/api/inscripciones/").json()["results"][0]
        self.assertEqual((fila["persona_foto"], fila["persona_sexo"]), (self.ana.foto_url, "F"))
        Usuario.objects.create(persona=self.ana, cuenta=self.admin)
        persona = self.client.get("/api/auth/sesion/").json()["persona"]
        self.assertEqual((persona["foto"], persona["sexo"]), (self.ana.foto_url, "F"))
        ficha = self.client.get(f"/api/personas/{self.ana.pk}/ficha/").json()["persona"]
        self.assertEqual(ficha["foto"], self.ana.foto_url)

    def test_quitarla(self):
        self.subir(self.archivo())
        r = self.client.delete(f"/api/personas/{self.ana.pk}/foto/")
        self.assertEqual(r.json(), {"foto": None})
        self.assertEqual(self.client.get(f"/api/personas/{self.ana.pk}/foto/").status_code, 404)

    def test_sin_permiso_de_cambiar_no_la_sube(self):
        from django.contrib.auth.models import Permission
        user = User.objects.create_user("lector", password="x")
        user.user_permissions.add(Permission.objects.get(codename="view_persona"))
        self.client.force_login(user)
        self.assertEqual(self.subir(self.archivo()).status_code, 403)

    def test_queda_en_la_auditoria_sin_la_imagen(self):
        from .models import Auditoria
        self.subir(self.archivo())
        registro = Auditoria.objects.get(modelo="persona", objeto_id=str(self.ana.pk),
                                         cambios__has_key="foto_actualizada")
        self.assertTrue(registro.cambios["foto_actualizada"]["despues"].startswith("Foto del "))
        self.assertFalse(Auditoria.objects.filter(modelo="fotopersona").exists())


class TableroPersonalizadoTests(Base):
    """Cada usuario elige qué secciones del tablero ver."""

    def patch(self, datos):
        return self.client.patch("/api/auth/preferencias/", datos,
                                 content_type="application/json")

    def test_empieza_mostrando_todo(self):
        pref = self.client.get("/api/auth/sesion/").json()["preferencias"]
        self.assertEqual(pref["tablero_ocultos"], [])

    def test_guarda_las_ocultas_sin_tocar_lo_demas(self):
        self.patch({"moneda": "USD"})
        r = self.patch({"tablero_ocultos": ["cumpleanos", "accesos", "cumpleanos"]})
        self.assertEqual(r.status_code, 200)
        pref = self.client.get("/api/auth/sesion/").json()["preferencias"]
        self.assertEqual(pref["tablero_ocultos"], ["cumpleanos", "accesos"])
        self.assertEqual(pref["moneda"], "USD")
        self.patch({"tablero_ocultos": []})
        pref = self.client.get("/api/auth/sesion/").json()["preferencias"]
        self.assertEqual(pref["tablero_ocultos"], [])

    def test_rechaza_lo_que_no_es_una_lista_de_claves(self):
        for malo in ["cumpleanos", [1], ["<script>"], ["x" * 41], ["a"] * 51]:
            self.assertEqual(self.patch({"tablero_ocultos": malo}).status_code, 400, malo)

    def test_es_de_cada_usuario(self):
        self.patch({"tablero_ocultos": ["cumpleanos"]})
        self.client.force_login(User.objects.create_superuser("otro", "o@example.com", "x"))
        pref = self.client.get("/api/auth/sesion/").json()["preferencias"]
        self.assertEqual(pref["tablero_ocultos"], [])


class CuentaServidorTests(Base):
    """La llave de la lista de Servidores: usuario, cuenta y rol en un paso."""

    def setUp(self):
        super().setUp()
        self.luis = self.persona("Luis Gómez", email="luis@example.com")
        self.servidor = Participacion.objects.create(
            retiro=self.retiro, persona=self.luis, tipo=Participacion.Tipo.SERVIDOR)
        self.ruta = f"/api/participaciones/{self.servidor.pk}/cuenta/"

    def test_crea_usuario_cuenta_y_rol_servidor_en_el_retiro(self):
        datos = self.client.get(self.ruta).json()
        self.assertEqual((datos["tiene_cuenta"], datos["usuario"], datos["email"]),
                         (False, None, "luis@example.com"))
        self.assertIn("Al guardar recibe el rol Servidor", datos["acceso"])
        r = self.post(self.ruta, {"username": "luis", "email": "luis@example.com",
                                  "password": "clave-larga-2027"})
        self.assertEqual(r.status_code, 200, r.content)
        self.assertIn("Tiene el rol Servidor", r.json()["acceso"])
        usuario = Usuario.objects.get(persona=self.luis)
        self.assertEqual(usuario.cuenta.username, "luis")
        asignacion = AsignacionRol.objects.get(usuario=usuario)
        self.assertEqual((asignacion.rol.codigo, asignacion.retiro_id),
                         (RolSistema.SERVIDOR, self.retiro.pk))
        # Cambiar la contraseña después no repite el rol.
        r = self.post(self.ruta, {"username": "luis", "password": "otra-clave-2027"})
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(AsignacionRol.objects.filter(usuario=usuario).count(), 1)
        self.client.logout()
        self.assertTrue(self.client.login(username="luis", password="otra-clave-2027"))

    def test_usa_el_usuario_que_ya_tenia(self):
        usuario = Usuario.objects.create(persona=self.luis)
        self.post(self.ruta, {"username": "luis", "password": "clave-larga-2027"})
        self.assertEqual(Usuario.objects.filter(persona=self.luis).get(), usuario)

    def test_con_errores_no_deja_nada_a_medias(self):
        r = self.post(self.ruta, {"username": "luis", "password": "corta"})
        self.assertIn("password", r.json())
        self.assertFalse(Usuario.objects.filter(persona=self.luis).exists())
        self.assertFalse(AsignacionRol.objects.exists())

    def test_no_es_para_caminantes(self):
        caminante = Participacion.objects.create(
            retiro=self.retiro, persona=self.persona("Pedro"),
            tipo=Participacion.Tipo.CAMINANTE)
        r = self.client.get(f"/api/participaciones/{caminante.pk}/cuenta/")
        self.assertEqual(r.status_code, 404)

    def entrar_con(self, codigo, **ambito):
        cuenta = User.objects.create_user(f"u{Usuario.objects.count()}", password="x")
        usuario = Usuario.objects.create(persona=self.persona(cuenta.username), cuenta=cuenta)
        AsignacionRol.objects.create(usuario=usuario, rol=rol_de(codigo), **ambito)
        self.client.force_login(cuenta)

    def sesion(self):
        return self.client.get("/api/auth/sesion/").json()

    def test_solo_super_admin_y_logistica_de_ese_retiro(self):
        self.assertTrue(self.sesion()["da_acceso_servidores"])  # superusuario
        self.entrar_con(RolSistema.SUPER_ADMIN)
        self.assertTrue(self.sesion()["da_acceso_servidores"])
        self.assertEqual(self.client.get(self.ruta).status_code, 200)

        self.entrar_con(RolSistema.LOGISTICA, retiro=self.retiro)
        self.assertTrue(self.sesion()["da_acceso_servidores"])
        r = self.post(self.ruta, {"username": "luis", "password": "clave-larga-2027"})
        self.assertEqual(r.status_code, 200, r.content)

        otro = Retiro.objects.create(
            parroquia=self.parroquia, tipo=self.retiro.tipo, edicion=13,
            fecha_inicio=date(2027, 9, 10), fecha_fin=date(2027, 9, 12))
        self.entrar_con(RolSistema.LOGISTICA, retiro=otro)
        self.assertEqual(self.client.get(self.ruta).status_code, 403)

        for codigo, ambito in [(RolSistema.COORD_PARROQUIA, {"parroquia": self.parroquia}),
                               (RolSistema.COORD_RETIRO, {"retiro": self.retiro}),
                               (RolSistema.TESORERIA, {"parroquia": self.parroquia})]:
            self.entrar_con(codigo, **ambito)
            self.assertFalse(self.sesion()["da_acceso_servidores"], codigo)
            self.assertEqual(self.client.get(self.ruta).status_code, 403, codigo)


class ContactosFamiliaresTests(Base):
    """El listado de familiares de contacto de caminantes y servidores."""

    def setUp(self):
        super().setUp()
        Inscripcion.objects.create(
            retiro=self.retiro, persona=self.persona("Ana Pérez", documento="12345678"),
            cama="7", contacto_emergencia={"nombre": "Rosa", "relacion": "Madre",
                                           "celular": "0424", "nombre3": "Luz",
                                           "telefono3": "0414"})
        Participacion.objects.create(
            retiro=self.retiro, persona=self.persona("Luis Gil"),
            tipo=Participacion.Tipo.SERVIDOR, duerme_en_casa=True, cama="B-2")
        self.ruta = f"/api/contactos/?retiro={self.retiro.pk}"

    def test_caminantes_y_servidores_con_sus_familiares(self):
        filas = self.client.get(self.ruta).json()["filas"]
        self.assertEqual(
            [(f["persona"], f["tipo"], f["documento"], f["cama"], f["contactos"]) for f in filas],
            [("Ana Pérez", "Caminante", "V-12345678", "7", [
                {"nombre": "Rosa", "relacion": "Madre", "telefono": "0424"},
                {"nombre": "Luz", "relacion": "", "telefono": "0414"},
            ]), ("Luis Gil", "Servidor", "", "B-2", [])],
        )

    def test_solo_para_quien_ve_datos_sensibles(self):
        user = User.objects.create_user("lector", password="x")
        user.user_permissions.add(*Permission.objects.filter(
            codename__in=["view_retiro", "view_inscripcion", "view_participacion"]))
        self.client.force_login(user)
        self.assertEqual(self.client.get(self.ruta).status_code, 403)

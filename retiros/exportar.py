"""
Descarga de una tabla en Excel (xlsx) o PDF: la usan los listados de cada
módulo (/api/<slug>/exportar/) y las tablas del panel de un retiro
(/api/exportar/).

Cada celda es un texto, un número, una fecha, un sí/no o un Dinero; en Excel
los números, fechas y montos quedan como tales (se pueden sumar y ordenar) y
en el PDF salen con el formato de la pantalla.
"""

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from django.template.loader import render_to_string
from django.utils import timezone
from django.utils.text import slugify

FORMATOS = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}
SIMBOLOS = {"VES": "Bs.S", "USD": "$", "EUR": "€"}
MAX_FILAS = 10_000


@dataclass
class Dinero:
    monto: float
    moneda: str


def _miles(valor, decimales):
    """1234567.5 -> "1.234.567,50" (como en Venezuela)."""
    texto = f"{valor:,.{decimales}f}"
    return texto.replace(",", "_").replace(".", ",").replace("_", ".")


def texto(valor):
    """Lo que se ve de una celda en el PDF."""
    if valor is None or valor == "":
        return "—"
    if isinstance(valor, Dinero):
        return f"{SIMBOLOS.get(valor.moneda, valor.moneda)} {_miles(valor.monto, 2)}"
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if isinstance(valor, datetime):
        return timezone.localtime(valor).strftime("%d/%m/%Y %H:%M")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, (int, float, Decimal)):
        return _miles(valor, 0 if float(valor).is_integer() else 2)
    return str(valor)


def _excel(titulo, subtitulo, columnas, filas):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    libro = Workbook()
    hoja = libro.active
    hoja.title = slugify(titulo)[:31] or "datos"
    hoja.append([titulo])
    hoja["A1"].font = Font(bold=True, size=14)
    hoja.append([subtitulo] if subtitulo else [])
    hoja.append([])
    hoja.append(columnas)
    for celda in hoja[4]:
        celda.font = Font(bold=True)
        celda.fill = PatternFill("solid", fgColor="F2F4F7")
        celda.alignment = Alignment(vertical="center")
    hoja.freeze_panes = "A5"

    anchos = [len(c) for c in columnas]
    for fila in filas:
        valores = []
        for i, valor in enumerate(fila):
            if isinstance(valor, Dinero):
                valores.append(round(valor.monto, 2))
            elif isinstance(valor, bool):
                valores.append("Sí" if valor else "No")
            elif isinstance(valor, datetime):
                valores.append(timezone.localtime(valor).replace(tzinfo=None))
            else:
                valores.append(valor)
            anchos[i] = max(anchos[i], len(texto(valor)))
        hoja.append(valores)
        for celda, valor in zip(hoja[hoja.max_row], fila):
            if isinstance(valor, Dinero):
                simbolo = SIMBOLOS.get(valor.moneda, valor.moneda)
                celda.number_format = f'"{simbolo}" #,##0.00'
            elif isinstance(valor, datetime):
                celda.number_format = "dd/mm/yyyy hh:mm"
            elif isinstance(valor, date):
                celda.number_format = "dd/mm/yyyy"
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = min(max(ancho + 2, 8), 60)
    hoja.auto_filter.ref = f"A4:{get_column_letter(len(columnas))}{max(hoja.max_row, 4)}"

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _pdf(request, titulo, subtitulo, columnas, filas):
    from weasyprint import HTML  # pesado de importar: solo al generar un PDF

    html = render_to_string("retiros/tabla_pdf.html", {
        "titulo": titulo,
        "subtitulo": subtitulo,
        "columnas": columnas,
        # Montos, cifras y fechas sin partirse en dos líneas.
        "filas": [[(texto(v), isinstance(v, (Dinero, int, float, Decimal, date))
                    and not isinstance(v, bool)) for v in fila] for fila in filas],
        "horizontal": len(columnas) > 6,
        "generado": timezone.localtime().strftime("%d/%m/%Y %H:%M"),
        "usuario": request.user.get_full_name() or request.user.get_username(),
    })
    return HTML(string=html, base_url=request.build_absolute_uri()).write_pdf()


def descargar(request, formato, titulo, columnas, filas, subtitulo=""):
    """La tabla como archivo para descargar (formato "xlsx" o "pdf")."""
    if formato == "xlsx":
        contenido = _excel(titulo, subtitulo, columnas, filas)
    else:
        contenido = _pdf(request, titulo, subtitulo, columnas, filas)
    nombre = f"{slugify(titulo) or 'tabla'}-{timezone.localdate():%Y-%m-%d}.{formato}"
    resp = HttpResponse(contenido, content_type=FORMATOS[formato])
    resp["Content-Disposition"] = f'attachment; filename="{nombre}"'
    return resp
